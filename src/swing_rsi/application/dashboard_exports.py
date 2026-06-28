from __future__ import annotations

import json
import math
import re
from collections.abc import Mapping
from datetime import date, datetime
from io import BytesIO
from numbers import Number
from pathlib import Path
from typing import Any, cast

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from swing_rsi.config import ProjectPaths

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_API_KEY_QUERY_RE = re.compile(r"(?i)(apikey=)[^&\s]+")
_FMP_KEY_ASSIGNMENT_RE = re.compile(r"(?i)(FMP_API_KEY\s*=\s*)[^,\s]+")
_SECRET_MARKERS = (
    "FMP_API_KEY",
    "apikey=",
    "super-secret",
    "secret-status-key",
    "raw authenticated",
)
_MISSING_EXPORT_LABEL = "Not available"


def _redact_text(value: str) -> str:
    redacted = _API_KEY_QUERY_RE.sub(r"\1[redacted]", value)
    redacted = _FMP_KEY_ASSIGNMENT_RE.sub(r"\1[redacted]", redacted)
    for marker in _SECRET_MARKERS:
        redacted = redacted.replace(marker, "[redacted]")
        redacted = redacted.replace(marker.upper(), "[redacted]")
        redacted = redacted.replace(marker.lower(), "[redacted]")
    return redacted


def _is_missing(value: object) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(cast(Any, value)))
    except (TypeError, ValueError):
        return False


def _is_number(value: object) -> bool:
    return isinstance(value, Number) and not isinstance(value, bool)


def _is_datetime_like(value: object) -> bool:
    return isinstance(value, datetime | date | pd.Timestamp)


def _normalize_special_text(value: str) -> str:
    text = _redact_text(value)
    normalized = text.strip()
    if normalized in {"inf", "+inf", "Inf", "+Inf", "Infinity", "+Infinity"}:
        return "∞"
    if normalized in {"-inf", "-Inf", "-Infinity"}:
        return "-∞"
    if normalized in {"nan", "NaN", "NaT", "None", ""}:
        return _MISSING_EXPORT_LABEL
    return text


def _clean_cell(value: object) -> object:
    if _is_missing(value):
        return None
    if isinstance(value, dict | list | tuple):
        return _redact_text(json.dumps(value, sort_keys=True, default=str))
    if isinstance(value, bytes):
        return _redact_text(value.decode("utf-8", errors="replace"))
    if isinstance(value, str):
        return _normalize_special_text(value)
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        item = cast(Any, value).item()
        return _clean_cell(item)
    if isinstance(value, float) and not math.isfinite(value):
        return "∞" if value > 0.0 else "-∞"
    return value


def _stringify_export_value(value: object) -> str:
    cleaned = _clean_cell(value)
    if cleaned is None:
        return _MISSING_EXPORT_LABEL
    if isinstance(cleaned, datetime):
        return cleaned.isoformat()
    if isinstance(cleaned, date):
        return cleaned.isoformat()
    if isinstance(cleaned, bool):
        return "True" if cleaned else "False"
    return _normalize_special_text(str(cleaned))


def _normalizable_values(series: pd.Series) -> list[object]:
    cleaned_values = [_clean_cell(value) for value in series.tolist()]
    return [value for value in cleaned_values if value is not None]


def _normalize_export_column(series: pd.Series) -> pd.Series:
    values = _normalizable_values(series)
    if not values:
        return pd.Series([None] * len(series), index=series.index, dtype="object")
    if all(isinstance(value, bool) for value in values):
        return series.map(_clean_cell).astype("boolean")
    if all(_is_number(value) for value in values):
        return pd.to_numeric(series.map(_clean_cell), errors="coerce")
    if all(_is_datetime_like(value) for value in values):
        return pd.to_datetime(series.map(_clean_cell), errors="coerce")
    return series.map(_stringify_export_value).astype("string")


def normalize_table_for_export(frame: pd.DataFrame) -> pd.DataFrame:
    """Return a redacted, Arrow-safe table while preserving genuine numeric/date columns."""
    normalized = frame.copy()
    normalized.columns = [_redact_text(str(column)) for column in normalized.columns]
    for column in normalized.columns:
        normalized[column] = _normalize_export_column(normalized[column])
    return normalized


def _ensure_safe_bytes(payload: bytes) -> bytes:
    for marker in _SECRET_MARKERS:
        encoded = marker.encode("utf-8")
        if encoded in payload or marker.upper().encode("utf-8") in payload:
            raise ValueError("Export payload contains a redacted secret marker")
    return payload


def to_csv_bytes(frame: pd.DataFrame) -> bytes:
    payload = normalize_table_for_export(frame).to_csv(index=False).encode("utf-8")
    return _ensure_safe_bytes(payload)


def _safe_sheet_name(name: str, existing: set[str]) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("_", name).strip() or "sheet"
    base = cleaned[:31]
    candidate = base
    index = 2
    while candidate in existing:
        suffix = f"_{index}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        index += 1
    existing.add(candidate)
    return candidate


def _is_percentage_column(name: object) -> bool:
    normalized = str(name).strip().lower()
    return any(
        token in normalized
        for token in (
            " rate",
            "_rate",
            "pct",
            "percent",
            "probability",
            "return",
            "drawdown",
            "mfe",
            "mae",
        )
    )


def _column_width(values: list[object]) -> float:
    width = max((len(str(value)) if value is not None else 0 for value in values), default=0)
    return float(min(max(width + 2, 10), 48))


def _excel_cell_value(value: object) -> object:
    if _is_missing(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return _excel_cell_value(cast(Any, value).item())
    if isinstance(value, float) and not math.isfinite(value):
        return "∞" if value > 0.0 else "-∞"
    return value


def _write_sheet(workbook: Workbook, name: str, frame: pd.DataFrame, existing: set[str]) -> None:
    worksheet = workbook.create_sheet(_safe_sheet_name(name, existing))
    cleaned = normalize_table_for_export(frame)
    headers = [str(column) for column in cleaned.columns]
    worksheet.append(headers)
    for row in cleaned.itertuples(index=False, name=None):
        worksheet.append([_excel_cell_value(value) for value in row])

    worksheet.freeze_panes = "A2"
    if headers:
        worksheet.auto_filter.ref = worksheet.dimensions

    for column_index, column_name in enumerate(headers, start=1):
        letter = get_column_letter(column_index)
        column_values = [column_name, *cleaned.iloc[:, column_index - 1].head(500).tolist()]
        worksheet.column_dimensions[letter].width = _column_width(column_values)
        is_percentage = _is_percentage_column(column_name)
        for cell in worksheet[letter][1:]:
            value = cell.value
            if isinstance(value, datetime | date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(value, int):
                cell.number_format = "#,##0"
            elif isinstance(value, float):
                cell.number_format = "0.00%" if is_percentage else "#,##0.0000"


def to_xlsx_bytes(
    sheets: Mapping[str, pd.DataFrame] | pd.DataFrame,
    *,
    sheet_name: str = "data",
) -> bytes:
    normalized: Mapping[str, pd.DataFrame] = (
        {sheet_name: sheets} if isinstance(sheets, pd.DataFrame) else sheets
    )
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    existing: set[str] = set()
    if not normalized:
        _write_sheet(workbook, "empty", pd.DataFrame(), existing)
    for name, frame in normalized.items():
        _write_sheet(workbook, name, frame, existing)
    output = BytesIO()
    workbook.save(output)
    return _ensure_safe_bytes(output.getvalue())


def save_xlsx_report(
    root: str | Path,
    filename: str,
    sheets: Mapping[str, pd.DataFrame] | pd.DataFrame,
    *,
    sheet_name: str = "data",
) -> Path:
    export_dir = ProjectPaths(Path(root)).reports / "dashboard_exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_filename = Path(filename).name
    if not safe_filename.lower().endswith(".xlsx"):
        safe_filename = f"{safe_filename}.xlsx"
    output = export_dir / safe_filename
    output.write_bytes(to_xlsx_bytes(sheets, sheet_name=sheet_name))
    return output
