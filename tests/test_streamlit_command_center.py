from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
import pytest

from swing_rsi.application.dashboard_exports import (
    normalize_table_for_export,
    save_xlsx_report,
    to_csv_bytes,
    to_xlsx_bytes,
)
from swing_rsi.application.dashboard_service import (
    CommandSpec,
    candidate_detail_url,
    command_specs,
    complete_engine_snapshot_frames,
    dashboard_startup_state,
    gate_audit_frame,
    model_edge_status_frame,
    model_registry_frame,
    regime_cache_detail_frames,
    run_dashboard_command,
    save_development_fmp_settings,
    scanner_results_frame,
    signal_board_frame,
    signal_board_metrics,
)
from swing_rsi.data.loader import save_ohlcv_csv
from swing_rsi.engine.gates import make_gate
from swing_rsi.engine.registry import RegisteredModel, register_model
from swing_rsi.engine.storage import engine_connection, initialize_engine_db

streamlit_testing = pytest.importorskip("streamlit.testing.v1")
AppTest = streamlit_testing.AppTest


def _assert_no_streamlit_exceptions(app: Any) -> None:
    assert not app.exception, [str(exception.value) for exception in app.exception]


def _sample_frame(rows: int = 40) -> pd.DataFrame:
    dates = pd.bdate_range("2026-04-01", periods=rows)
    close = pd.Series(range(rows), index=dates, dtype=float) + 100.0
    return pd.DataFrame(
        {
            "Open": close - 0.2,
            "High": close + 1.0,
            "Low": close - 1.0,
            "Close": close,
            "Volume": 1_000_000.0,
        },
        index=dates.rename("Date"),
    )


@pytest.fixture()
def command_center_root(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    (tmp_path / "src" / "swing_rsi").mkdir(parents=True)
    (tmp_path / "pyproject.toml").write_text("[project]\n", encoding="utf-8")
    raw = tmp_path / "data" / "raw"
    raw.mkdir(parents=True)
    save_ohlcv_csv(_sample_frame(), raw / "DEMO.csv")
    (tmp_path / "data" / "manifests").mkdir(parents=True)
    (tmp_path / "data" / "manifests" / "DEMO.json").write_text(
        json.dumps(
            {
                "symbol": "DEMO",
                "provider": "fmp",
                "actual_last_date": "2026-05-26",
                "row_count": 40,
                "stale_data_status": "fresh_within_3_calendar_days",
                "raw_file_hash": "rawhash",
                "retrieval_timestamp_utc": "2026-05-26T20:00:00+00:00",
            }
        ),
        encoding="utf-8",
    )
    config_dir = tmp_path / "configs" / "universe"
    config_dir.mkdir(parents=True)
    (config_dir / "core.yaml").write_text(
        """
name: command-center-test
provider: fmp
default_start: "2020-01-01"
symbols:
  - symbol: DEMO
    enabled: true
    role: stock
    sector: test
relationships: []
""",
        encoding="utf-8",
    )
    artifact = tmp_path / "artifacts" / "models" / "model.joblib"
    artifact.parent.mkdir(parents=True)
    artifact.write_text("artifact", encoding="utf-8")
    db = tmp_path / "state" / "engine.sqlite3"
    initialize_engine_db(db)
    gate = make_gate(
        gate_id="profit_factor_min_090",
        gate_name="Profit Factor Minimum",
        category="performance",
        scope="development_holdout",
        metric_name="holdout_profit_factor",
        threshold=0.9,
        comparator=">=",
        actual_value=0.5,
        status="FAIL",
        mandatory=True,
        evidence_source="test",
        reason="Profit factor is below the minimum.",
        configuration_hash_value="policyhash",
    )
    register_model(
        db,
        RegisteredModel(
            model_id="model-demo",
            task="probability_positive_return",
            horizon=10,
            direction="bull",
            family="hist_gradient_boosting",
            state="CANDIDATE",
            training_start="2020-01-01",
            training_end="2021-01-01",
            validation_start="2021-01-04",
            validation_end="2021-06-01",
            holdout_start="2021-06-02",
            holdout_end="2022-01-01",
            universe_snapshot_id="universe",
            feature_manifest_hash="featurehash",
            raw_manifest_hashes=("rawhash",),
            hyperparameters={},
            metrics={
                "product_class_scope": "POOLED",
                "holdout_status": "DEVELOPMENT_HOLDOUT",
                "selected_holdout_samples": 3,
                "selected_observation_rate": 0.1,
                "holdout_mean_return_lcb_90": -0.01,
                "portfolio_max_drawdown": -0.02,
                "portfolio_total_return": 0.03,
            },
            calibration_metrics={"brier_skill_score": -0.1, "holdout_brier": 0.24},
            quality_gates={"profit_factor_min_090": False},
            artifact_path=str(artifact),
            code_commit_hash="abcdef",
            created_at_utc="2026-06-27T12:00:00+00:00",
            gate_results=(gate,),
        ),
    )
    scanner_dir = tmp_path / "artifacts" / "scanner"
    scanner_dir.mkdir(parents=True)
    scanner_path = scanner_dir / "scan-demo_scanner.csv"
    rejected_row = {
        "scan_id": "scan-demo",
        "as_of_date": "2026-05-26",
        "ticker": "DEMO",
        "product_class_scope": "POOLED",
        "direction": "Bullish",
        "horizon": 10,
        "model_id": "model-demo",
        "calibrated_probability": 0.61,
        "expected_return": 0.02,
        "expected_mfe": 0.04,
        "expected_mae": -0.01,
        "target_before_stop_probability": 0.55,
        "candidate_status": "REJECTED",
        "exclusion_reason": "model_not_promoted; failed_gate",
        "top_attribution_categories": "Model contribution",
        "supporting_evidence": "Supporting evidence",
        "historical_analogs": "[]",
        "feature_snapshot_hash": "featurehash",
    }
    shadow_row = {
        "scan_id": "scan-demo",
        "as_of_date": "2026-05-26",
        "ticker": "DEMO2",
        "product_class_scope": "POOLED",
        "direction": "Bullish",
        "horizon": 10,
        "model_id": "model-demo",
        "calibrated_probability": 0.63,
        "expected_return": 0.03,
        "expected_mfe": None,
        "expected_mae": None,
        "target_before_stop_probability": 0.57,
        "candidate_status": "ACTIONABLE_PAPER_CANDIDATE",
        "exclusion_reason": "",
        "top_attribution_categories": "Model contribution",
        "supporting_evidence": "Supporting evidence",
        "historical_analogs": "[]",
        "feature_snapshot_hash": "featurehash",
    }
    pd.DataFrame(
        [
            rejected_row,
            shadow_row,
        ]
    ).to_csv(scanner_path, index=False)
    with engine_connection(db) as connection:
        connection.execute(
            """
            INSERT INTO scanner_snapshots VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "scan-demo",
                "2026-05-26",
                "2026-06-27T12:30:00+00:00",
                json.dumps(["model-demo"]),
                "universe",
                "featurehash",
                str(scanner_path),
                str(scanner_path.with_suffix(".parquet")),
                2,
                json.dumps({"rejected_rows": 1, "actionable_rows": 1}),
            ),
        )
        connection.execute(
            """
            INSERT INTO final_holdout_runs (
                run_id, schema_version, created_at_utc, creation_git_commit,
                baseline_market_date, first_eligible_future_signal_date, universe_snapshot_id,
                feature_manifest_hash, generation_id, model_ids_json, scanner_identity_version,
                execution_policy_hash, sample_policy_version, sample_policy_hash,
                sample_policy_json, horizon, direction, status, invalidation_reason,
                latest_processed_market_date, metadata_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "run-demo",
                "prospective_final_holdout_v1",
                "2026-06-27T12:40:00+00:00",
                "abcdef",
                "2026-06-25",
                "2026-06-26",
                "universe",
                "featurehash",
                "2026-06-27T12:00:00+00:00",
                json.dumps(["model-demo"]),
                1,
                "executionhash",
                "prospective_final_holdout_sample_v1",
                "samplehash",
                json.dumps({}),
                10,
                "bull",
                "COLLECTING",
                None,
                "2026-06-26",
                json.dumps({"processed_sessions": ["2026-06-26"]}),
            ),
        )
        connection.execute(
            """
            INSERT INTO final_holdout_models VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "run-demo",
                "model-demo",
                "2026-06-27T12:00:00+00:00",
                str(artifact),
                "bad-hash",
                "CANDIDATE",
                0,
                1,
                "selectionhash",
                "calhash",
                "oodhash",
                json.dumps(["development_gate_failed"]),
                json.dumps({"feature_manifest_hash": "featurehash"}),
            ),
        )
        shadow_payload = {
            "run_id": "run-demo",
            "mode": "SHADOW_FINAL_HOLDOUT",
            "not_live_trade_recommendation": True,
            "scanner_row": shadow_row,
            "entry_rule": "next_completed_session_open",
            "planned_entry_date": "2026-05-27",
            "expected_return": shadow_row["expected_return"],
            "expected_mfe": shadow_row["expected_mfe"],
            "expected_mae": shadow_row["expected_mae"],
            "target_before_stop_probability": shadow_row["target_before_stop_probability"],
        }
        for event_id, unique_key, event_type in (
            (
                "signal-demo2-event",
                "signal|scan-demo|DEMO2|Bullish|model-demo",
                "FINAL_HOLDOUT_SIGNAL_CREATED",
            ),
            (
                "pending-demo2-event",
                "pending|scan-demo|DEMO2|Bullish|model-demo",
                "FINAL_HOLDOUT_ENTRY_PENDING",
            ),
        ):
            connection.execute(
                """
                INSERT INTO forward_events VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    event_id,
                    unique_key,
                    event_type,
                    "2026-06-27T12:45:00+00:00",
                    "2026-05-26",
                    "DEMO2",
                    "Bullish",
                    "model-demo",
                    "scan-demo",
                    "featurehash",
                    json.dumps(shadow_payload),
                ),
            )
    monkeypatch.setattr("dashboard.ui.components.resolve_project_root", lambda _: tmp_path)

    def fail_download(*_: object, **__: object) -> None:
        raise AssertionError("Command Center page load attempted an FMP download")

    monkeypatch.setattr("swing_rsi.data.loader.download_daily", fail_download)
    monkeypatch.delenv("FMP_API_KEY", raising=False)
    return tmp_path


def _write_regime_cache_metadata(
    root: Path,
    *,
    include_latest_run_fields: bool = True,
    reason: str = "cache_valid",
) -> tuple[Path, Path]:
    cache_dir = root / "data" / "cache" / "regime"
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / "universe-cache_expanding_kmeans_regime_cache_v1.json"
    status_path = cache_dir / "universe-cache_expanding_kmeans_regime_cache_status_v1.json"
    cache_path.write_text(
        json.dumps(
            {
                "schema_version": "expanding_kmeans_regime_cache_v1",
                "universe_snapshot_id": "universe-cache",
                "feature_manifest_hash": (
                    "3007abe80b54d40c87566bc9185a3086098c91b95d156c664ced290242090ae5"
                ),
                "regime_input_columns": [
                    "market_regime_trend_score",
                    "market_regime_volatility_score",
                    "breadth_advance_pct",
                    "breadth_dispersion_20",
                ],
                "kmeans_parameters": {
                    "n_clusters": 3,
                    "n_init": 10,
                    "random_state": 42,
                },
                "random_seed": 42,
                "scaler_preprocessing_configuration": {"scaler": "none"},
                "minimum_sample_requirement": 126,
                "input_prefix_hash": "input-prefix-hash",
                "full_input_hash": "full-input-hash",
                "output_dataframe_hash": "output-dataframe-hash",
                "dates_covered": ["2026-06-24", "2026-06-25", "2026-06-26"],
                "last_cached_date": "2026-06-26",
                "row_count": 5005,
                "symbols_covered": ["AAPL", "SPY"],
                "output_column_names": ["Date", "market_regime_cluster_expanding"],
                "creation_timestamp_utc": "2026-06-28T00:00:00+00:00",
                "update_timestamp_utc": "2026-06-28T00:01:00+00:00",
                "cache_validity_status": "VALID",
                "records": [],
            },
            sort_keys=True,
        ),
        encoding="utf-8",
    )
    status_payload: dict[str, object] = {
        "schema_version": "expanding_kmeans_regime_cache_status_v1",
        "cache_schema_version": "expanding_kmeans_regime_cache_v1",
        "universe_snapshot_id": "universe-cache",
        "feature_manifest_hash": "3007abe80b54d40c87566bc9185a3086098c91b95d156c664ced290242090ae5",
        "feature_builder_version": "features.py:_expanding_kmeans_regime:v1",
        "status": "HIT",
        "reason": reason,
        "cache_validity_status": "VALID",
        "cache_path": str(cache_path),
        "last_cached_date": "2026-06-26",
        "row_count": 5005,
        "update_timestamp_utc": "2026-06-28T00:02:00+00:00",
    }
    if include_latest_run_fields:
        status_payload.update(
            {
                "cached_dates_reused": 5005,
                "new_dates_computed": 0,
                "kmeans_fits_avoided": 2518,
                "kmeans_fits_performed": 0,
                "regime_runtime_seconds": 2.5,
            }
        )
    status_path.write_text(json.dumps(status_payload, sort_keys=True), encoding="utf-8")
    return cache_path, status_path


def test_dashboard_blocks_startup_from_operational_repo(monkeypatch: pytest.MonkeyPatch) -> None:
    from swing_rsi.application import dashboard_service as service

    monkeypatch.setattr(service, "OPERATIONAL_REPOSITORY", Path("/tmp/ops"))
    monkeypatch.setattr(service, "DEVELOPMENT_REPOSITORY", Path("/tmp/dev"))
    state = dashboard_startup_state("/tmp/ops")
    assert state.is_operational_repository is True
    assert "Blocked" in state.message


def test_dashboard_identifies_development_repo(monkeypatch: pytest.MonkeyPatch) -> None:
    from swing_rsi.application import dashboard_service as service

    monkeypatch.setattr(service, "DEVELOPMENT_REPOSITORY", Path("/tmp/dev"))
    monkeypatch.setattr(service, "OPERATIONAL_REPOSITORY", Path("/tmp/ops"))
    state = dashboard_startup_state("/tmp/dev")
    assert state.is_development_worktree is True
    assert state.requires_confirmation is False


def test_pages_load_without_fmp_key_or_mutation(
    command_center_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    db = command_center_root / "state" / "engine.sqlite3"
    scanner = command_center_root / "artifacts" / "scanner" / "scan-demo_scanner.csv"
    artifact = command_center_root / "artifacts" / "models" / "model.joblib"
    before_db = db.read_bytes()
    before_scanner = scanner.read_bytes()
    before_artifact = artifact.read_bytes()

    def fail_build_features(*_: object, **__: object) -> None:
        raise AssertionError("Dashboard page load attempted to run build-features")

    monkeypatch.setattr(
        "swing_rsi.application.engine_service.build_autonomous_features",
        fail_build_features,
    )
    for page in (
        "dashboard/sections/overview.py",
        "dashboard/sections/signal_board.py",
        "dashboard/sections/shadow_forward_test.py",
        "dashboard/sections/model_edge_status.py",
        "dashboard/sections/scanner_results.py",
        "dashboard/sections/candidate_detail.py",
        "dashboard/sections/product_class_research.py",
        "dashboard/sections/gate_audit.py",
        "dashboard/sections/data_universe.py",
        "dashboard/sections/reports_and_exports.py",
        "dashboard/sections/engine_commands.py",
        "dashboard/sections/baselines_legacy.py",
        "dashboard/sections/developer_diagnostics.py",
    ):
        app = AppTest.from_file(page).run(timeout=30)
        _assert_no_streamlit_exceptions(app)
    assert db.read_bytes() == before_db
    assert scanner.read_bytes() == before_scanner
    assert artifact.read_bytes() == before_artifact


def test_overview_displays_regime_cache_status(command_center_root: Path) -> None:
    _write_regime_cache_metadata(command_center_root)

    app = AppTest.from_file("dashboard/sections/overview.py").run(timeout=30)

    _assert_no_streamlit_exceptions(app)
    assert "Regime KMeans Cache" in {subheader.value for subheader in app.subheader}
    metrics = {metric.label: metric.value for metric in app.metric}
    assert metrics["Regime cache status"] == "HIT"
    assert metrics["Last cached date"] == "2026-06-26"
    assert metrics["KMeans fits avoided"] == "2,518"
    assert metrics["Regime runtime"] == "2.50 seconds"
    assert metrics["Cache validity reason"] == "cache_valid"


def test_developer_diagnostics_displays_regime_cache_section(
    command_center_root: Path,
) -> None:
    _write_regime_cache_metadata(command_center_root)

    app = AppTest.from_file("dashboard/sections/developer_diagnostics.py").run(timeout=30)

    _assert_no_streamlit_exceptions(app)
    assert "Regime KMeans Cache" in {subheader.value for subheader in app.subheader}
    metrics = {metric.label: metric.value for metric in app.metric}
    assert metrics["Regime cache status"] == "HIT"
    assert metrics["KMeans fits avoided"] == "2,518"
    assert metrics["KMeans fits performed"] == "0"
    assert metrics["Regime runtime"] == "2.50 seconds"
    assert any("Nothing required. build-features is using" in info.value for info in app.info)


def test_missing_regime_cache_metadata_displays_not_found(
    command_center_root: Path,
) -> None:
    frames = regime_cache_detail_frames(command_center_root)
    summary = frames["summary"].iloc[0]

    assert summary["status"] == "NOT_FOUND"
    assert summary["validity_reason"] == "cache_missing"

    app = AppTest.from_file("dashboard/sections/overview.py").run(timeout=30)
    _assert_no_streamlit_exceptions(app)
    assert any("Regime cache: Not found" in info.value for info in app.info)


def test_regime_cache_missing_fields_display_not_available(command_center_root: Path) -> None:
    _write_regime_cache_metadata(command_center_root, include_latest_run_fields=False)

    summary = regime_cache_detail_frames(command_center_root)["summary"].iloc[0]

    assert summary["status"] == "HIT"
    assert summary["validity_reason"] == "cache_valid"
    assert summary["kmeans_fits_avoided"] == "Not available"
    assert summary["kmeans_fits_performed"] == "Not available"
    assert summary["regime_runtime_seconds"] == "Not available"


def test_regime_cache_exports_csv_and_xlsx_without_secrets(command_center_root: Path) -> None:
    _write_regime_cache_metadata(
        command_center_root,
        reason="cache_valid; apikey=secret-token; FMP_API_KEY=secret-token",
    )
    frames = regime_cache_detail_frames(command_center_root)

    csv_bytes = to_csv_bytes(frames["metadata"])
    xlsx_bytes = to_xlsx_bytes(frames)
    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes))

    assert csv_bytes.startswith(b"field,value")
    assert set(workbook.sheetnames) == {"summary", "metadata", "input_columns", "kmeans_config"}
    assert b"secret-token" not in csv_bytes
    assert b"secret-token" not in xlsx_bytes
    assert b"FMP_API_KEY" not in csv_bytes
    assert b"apikey=secret-token" not in csv_bytes


def test_app_shell_exposes_password_only_fmp_settings(command_center_root: Path) -> None:
    app = AppTest.from_file("dashboard/app.py").run(timeout=30)
    _assert_no_streamlit_exceptions(app)
    labels = {widget.label for widget in app.text_input}
    assert "FMP API key" in labels
    source = Path("dashboard/ui/components.py").read_text(encoding="utf-8")
    assert 'streamlit.text_input("FMP API key", type="password", value="")' in source


def test_model_registry_and_gate_audit_load_local_state(command_center_root: Path) -> None:
    models = model_registry_frame(command_center_root)
    gates = gate_audit_frame(command_center_root)
    edges = model_edge_status_frame(command_center_root)
    assert "model-demo" in set(models["model_id"])
    assert "FAIL" in set(gates["status"])
    assert "SHADOW VALIDATION" in set(edges["edge_status"])


def test_signal_first_tables_label_live_shadow_and_missing_paths(
    command_center_root: Path,
) -> None:
    metrics = signal_board_metrics(command_center_root)
    board = signal_board_frame(command_center_root)
    scanner = scanner_results_frame(command_center_root)

    assert metrics["live_actionable_signals"] == 0
    assert metrics["shadow_paper_signals"] == 1
    assert metrics["pending_entries"] == 1
    assert "REJECTED" in set(board["live_shadow_rejected_classification"])
    assert "SHADOW ONLY" in set(board["live_shadow_rejected_classification"])
    shadow = board.loc[board["ticker"].astype(str) == "DEMO2"].iloc[0]
    assert shadow["signal_status"] == "PENDING ENTRY"
    assert shadow["expected_mfe"] == "Not available"
    assert shadow["expected_mae"] == "Not available"
    assert shadow["next_required_event"] == "Pending entry waits for next session open."
    assert (
        shadow["why_shadow_only"] == "This row is waiting for the next eligible session open. "
        "It is not a live trade recommendation."
    )
    joined = " ".join(scanner.astype(str).stack().tolist())
    assert "Trade signal" not in joined
    rejected = scanner.loc[scanner["candidate_classification"].astype(str) == "REJECTED"]
    assert not rejected.empty


def test_signal_board_default_sections_keep_shadow_and_pending_visible(
    command_center_root: Path,
) -> None:
    db = command_center_root / "state" / "engine.sqlite3"
    before_db = db.read_bytes()

    app = AppTest.from_file("dashboard/sections/signal_board.py").run(timeout=30)

    _assert_no_streamlit_exceptions(app)
    assert db.read_bytes() == before_db
    assert app.toggle[0].label == "Show rejected/research rows"
    assert app.toggle[0].value is False
    subheaders = {subheader.value for subheader in app.subheader}
    assert "Live Actionable Signals: 0 rows" in subheaders
    assert "Shadow / Paper Signals: 1 row" in subheaders
    assert "Pending Entries: 1 row" in subheaders
    assert "Open Shadow Positions: 0 rows" in subheaders
    assert "Closed / Matured Outcomes: 0 rows" in subheaders
    assert any(
        "No live actionable signals. Shadow and pending rows remain visible below." in info.value
        for info in app.info
    )
    assert any(
        "This row is waiting for the next eligible session open. "
        "It is not a live trade recommendation." in info.value
        for info in app.info
    )


def test_signal_board_detail_links_preselect_candidate_detail(
    command_center_root: Path,
) -> None:
    board = signal_board_frame(command_center_root)
    linked = board.loc[board["ticker"].astype(str) == "DEMO2"].iloc[0]

    assert linked["candidate_detail_url"] == candidate_detail_url(
        scan_id="scan-demo",
        ticker="DEMO2",
        model_id="model-demo",
        direction="Bullish",
        run_id="run-demo",
        event_id="pending-demo2-event",
        status="PENDING ENTRY",
    )
    assert linked["open_url"] == linked["candidate_detail_url"]
    assert linked["candidate_detail_url"].startswith("/candidate-detail?")
    assert "run_id=run-demo" in linked["open_url"]
    assert "event_id=pending-demo2-event" in linked["open_url"]
    assert "model_id=model-demo" in linked["open_url"]
    assert "ticker=DEMO2" in linked["open_url"]
    assert "direction=Bullish" in linked["open_url"]

    app = AppTest.from_file("dashboard/sections/candidate_detail.py")
    app.query_params["scan_id"] = "scan-demo"
    app.query_params["ticker"] = "DEMO2"
    app.query_params["model_id"] = "model-demo"
    app.query_params["direction"] = "Bullish"
    app.run(timeout=30)

    _assert_no_streamlit_exceptions(app)
    assert app.selectbox[0].label == "Scan ID"
    assert app.selectbox[0].value == "scan-demo"
    assert app.selectbox[1].label == "Ticker"
    assert app.selectbox[1].value == "DEMO2"
    assert app.selectbox[2].label == "Model ID"
    assert app.selectbox[2].value == "model-demo"
    assert any(success.value == "Opened from Signal Board selection." for success in app.success)


def test_complete_engine_snapshot_export_contains_required_sheets(
    command_center_root: Path,
    tmp_path: Path,
) -> None:
    frames = complete_engine_snapshot_frames(command_center_root)
    xlsx_bytes = to_xlsx_bytes(frames)
    workbook_path = tmp_path / "complete_snapshot.xlsx"
    workbook_path.write_bytes(xlsx_bytes)
    workbook = openpyxl.load_workbook(workbook_path)
    assert set(workbook.sheetnames) == {
        "signal_board",
        "shadow_forward_status",
        "model_edge_status",
        "scanner_results",
        "gate_audit",
        "product_class_research",
        "candidate_attribution",
        "data_universe",
        "reports_index",
    }
    assert workbook["signal_board"]["A1"].value == "ticker"
    assert "secret" not in workbook_path.read_bytes().decode("latin1", errors="ignore")


def test_exports_create_csv_and_xlsx_without_secrets(tmp_path: Path) -> None:
    frame = pd.DataFrame([{"model_id": "m1", "return": 0.12, "note": "apikey=secret"}])
    csv_bytes = to_csv_bytes(frame)
    assert b"apikey=secret" not in csv_bytes
    xlsx_bytes = to_xlsx_bytes({"models": frame, "gates": frame})
    workbook_path = tmp_path / "test.xlsx"
    workbook_path.write_bytes(xlsx_bytes)
    workbook = openpyxl.load_workbook(workbook_path)
    assert set(workbook.sheetnames) == {"models", "gates"}
    assert workbook["models"]["A1"].value == "model_id"
    assert "secret" not in workbook_path.read_bytes().decode("latin1", errors="ignore")


def test_mixed_value_export_normalization_is_arrow_safe(tmp_path: Path) -> None:
    frame = pd.DataFrame(
        {
            "Metric": [
                "text",
                "int",
                "float",
                "bool",
                "none",
                "nan",
                "positive_inf",
                "negative_inf",
                "secret",
            ],
            "Value": [
                "alpha",
                np.int64(7),
                3.25,
                True,
                None,
                np.nan,
                np.inf,
                -np.inf,
                "FMP_API_KEY=super-secret",
            ],
        }
    )

    normalized = normalize_table_for_export(frame)

    assert str(normalized["Value"].dtype) == "string"
    values = normalized["Value"].tolist()
    assert "0" not in values
    assert "Not available" in values
    assert "∞" in values
    assert "-∞" in values
    assert "7" in values
    assert "3.25" in values
    assert "True" in values
    assert not any("FMP_API_KEY" in value or "super-secret" in value for value in values)

    csv_bytes = to_csv_bytes(frame)
    assert b"Not available" in csv_bytes
    assert b"FMP_API_KEY" not in csv_bytes
    assert b"super-secret" not in csv_bytes

    xlsx_bytes = to_xlsx_bytes({"mixed": frame})
    workbook_path = tmp_path / "mixed.xlsx"
    workbook_path.write_bytes(xlsx_bytes)
    workbook = openpyxl.load_workbook(workbook_path)
    sheet_values = [cell.value for cell in workbook["mixed"]["B"][1:]]
    assert "Not available" in sheet_values
    assert "∞" in sheet_values
    assert "-∞" in sheet_values

    parquet_buffer = BytesIO()
    normalized.to_parquet(parquet_buffer, index=False)
    parquet_buffer.seek(0)
    round_trip = pd.read_parquet(parquet_buffer)
    assert round_trip["Value"].tolist() == values


def test_export_normalization_preserves_numeric_tables() -> None:
    frame = pd.DataFrame(
        {
            "metric": ["a", "b", "c"],
            "value": [1, 2, 3],
            "rate": [0.1, np.nan, 0.3],
        }
    )

    normalized = normalize_table_for_export(frame)

    assert pd.api.types.is_integer_dtype(normalized["value"])
    assert pd.api.types.is_float_dtype(normalized["rate"])
    assert pd.isna(normalized.loc[1, "rate"])


def test_save_xlsx_report_uses_ignored_dashboard_export_directory(tmp_path: Path) -> None:
    output = save_xlsx_report(tmp_path, "sample.xlsx", pd.DataFrame([{"a": 1}]))
    assert output.relative_to(tmp_path).as_posix() == "reports/dashboard_exports/sample.xlsx"


def test_dashboard_export_directory_is_git_ignored() -> None:
    ignore = Path(".gitignore").read_text(encoding="utf-8")
    assert "reports/dashboard_exports/" in ignore
    assert "reports/dashboard_command_logs/" in ignore


def test_engine_commands_exclude_discovery_and_promotion(command_center_root: Path) -> None:
    labels = {spec.label for spec in command_specs(fmp_configured=False)}
    assert "discover-models" not in labels
    assert "promote-model" not in labels
    assert "universe-update" not in labels


def test_engine_command_buttons_require_confirmation_and_disable_mutations(
    command_center_root: Path,
) -> None:
    app = AppTest.from_file("dashboard/sections/engine_commands.py").run(timeout=30)
    _assert_no_streamlit_exceptions(app)

    buttons = {button.label: button for button in app.button}
    assert buttons["Run selected command"].disabled is True
    for label in (
        "discover-models disabled",
        "promote-model disabled",
        "final-holdout-init disabled",
        "forward-update disabled",
    ):
        assert buttons[label].disabled is True

    app.checkbox[0].set_value(True)
    app.run(timeout=30)
    _assert_no_streamlit_exceptions(app)
    buttons = {button.label: button for button in app.button}
    assert buttons["Run selected command"].disabled is False


def test_command_runner_refuses_operational_repo(monkeypatch: pytest.MonkeyPatch) -> None:
    from swing_rsi.application import dashboard_service as service

    monkeypatch.setattr(service, "OPERATIONAL_REPOSITORY", Path("/tmp/ops"))
    monkeypatch.setattr(service, "DEVELOPMENT_REPOSITORY", Path("/tmp/dev"))
    spec = CommandSpec("test", "test", (sys.executable, "-c", "print('ok')"), mutates=False)
    with pytest.raises(ValueError, match="operational"):
        run_dashboard_command("/tmp/ops", spec)


def test_command_runner_captures_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from swing_rsi.application import dashboard_service as service

    (tmp_path / "reports").mkdir()
    monkeypatch.setattr(service, "DEVELOPMENT_REPOSITORY", tmp_path)
    monkeypatch.setattr(service, "OPERATIONAL_REPOSITORY", tmp_path / "ops")
    spec = CommandSpec("test", "test", (sys.executable, "-c", "print('ok')"), mutates=False)
    result = run_dashboard_command(tmp_path, spec)
    assert result.return_code == 0
    assert "ok" in result.stdout
    assert result.log_path.exists()


def test_fmp_settings_button_backend_writes_dev_env_without_echoing_key(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from swing_rsi.application import dashboard_service as service

    monkeypatch.setattr(service, "DEVELOPMENT_REPOSITORY", tmp_path)
    monkeypatch.setattr(service, "OPERATIONAL_REPOSITORY", tmp_path / "ops")
    result = save_development_fmp_settings(root=tmp_path, api_key="abc123")
    assert result.fmp_configured is True
    assert "abc123" not in str(result)
    assert (tmp_path / ".env").exists()
    monkeypatch.delenv("FMP_API_KEY", raising=False)
    monkeypatch.delenv("FMP_BASE_URL", raising=False)
