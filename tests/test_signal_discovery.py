from __future__ import annotations

import json
import math
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from swing_rsi.application.dashboard_exports import to_xlsx_bytes
from swing_rsi.engine.features import numeric_feature_columns, reject_label_columns
from swing_rsi.engine.labels import LabelConfig, build_symbol_labels
from swing_rsi.engine.signal_discovery import (
    CALIBRATION_DIAGNOSTIC_THRESHOLDS,
    MULTI_ANGLE_CALIBRATION_AUDIT_SCHEMA_VERSION,
    SIGNAL_DISCOVERY_GENERATION_TYPE,
    SIGNAL_DISCOVERY_SCHEMA_VERSION,
    default_archetype_registry,
    default_hypothesis_registry,
    export_signal_discovery_blocker_report,
    export_signal_discovery_generation,
    historical_analog_robustness_frames_from_analogs,
    load_signal_discovery_frames,
    run_signal_discovery,
    signal_discovery_analog_robustness_frames,
    signal_discovery_blocked_analog_frames,
    signal_discovery_blocker_report_frames,
)
from swing_rsi.engine.target_stop_policy import (
    SECTOR_ROTATION_BUY_ORDINARY_CALIBRATION_EVIDENCE,
    SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_HYPOTHESIS_ID,
    SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID,
    CalibrationPolicyEvidence,
    augment_model_frame_with_policy_outcomes,
    candidate_policy_outcome_labels,
    select_sector_rotation_buy_ordinary_policy_candidate,
    target_stop_policy_registry,
)
from swing_rsi.engine.time_exit_utility import (
    SECTOR_ROTATION_BUY_ORDINARY_TIME_EXIT_HYPOTHESIS_ID,
    TIME_EXIT_BASELINE_POLICY_ALIAS,
    TIME_EXIT_UTILITY_LABEL_SCHEMA_VERSION,
    augment_model_frame_with_time_exit_utility_labels,
    build_time_exit_utility_labels,
    time_exit_quality_bucket,
    time_exit_utility_calibration_summary_frame,
    time_exit_utility_outcome_labels,
)


def _write_universe(root: Path) -> None:
    config_dir = root / "configs" / "universe"
    config_dir.mkdir(parents=True, exist_ok=True)
    (config_dir / "core.yaml").write_text(
        """
name: signal-discovery-test
provider: fmp
default_start: "2020-01-01"
symbols:
  - symbol: AAA
    enabled: true
    role: stock
    sector: test
  - symbol: BBB
    enabled: true
    role: stock
    sector: test
  - symbol: CCC
    enabled: true
    role: ordinary_etf
    sector: test
  - symbol: TZA
    enabled: true
    role: leveraged_inverse_etf
relationships:
  - source: IWM
    related: [TZA]
""",
        encoding="utf-8",
    )


def _write_config(
    root: Path,
    *,
    signal_score_threshold: float = 0.10,
    probability_threshold: float = 0.10,
    candidate_cap: int = 1,
) -> Path:
    config_dir = root / "configs" / "signal_discovery"
    config_dir.mkdir(parents=True, exist_ok=True)
    path = config_dir / "test_v1.yaml"
    path.write_text(
        f"""
schema_version: multi_angle_signal_discovery_v1
enabled_hypotheses:
  - reversal_buy_5d
  - reversal_sell_5d
horizons: [5]
product_scopes: [POOLED, ORDINARY, LEVERAGED_INVERSE]
model_families:
  - hist_gradient_boosting
  - extra_trees
minimum_training_samples: 120
minimum_calibration_samples: 24
minimum_holdout_samples: 24
max_selected_features: 10
analog_count: 3
candidate_cap_per_hypothesis: {candidate_cap}
random_seed: 17
research_start: "2026-01-01"
research_end:
costs:
  round_trip_bps: 5.0
selection_policy:
  reference: signal_discovery_policy_test_v1
  probability: {probability_threshold}
  target_before_stop: 0.10
  expected_return: -0.10
  signal_score: {signal_score_threshold}
ood_policy:
  reference: prediction_ood_governance_test_v1
  feature_rate_limit: 1.00
""",
        encoding="utf-8",
    )
    return path


def _write_feature_data(root: Path, *, db_marker: bool = False) -> None:
    feature_dir = root / "data" / "features"
    feature_dir.mkdir(parents=True, exist_ok=True)
    manifest_dir = root / "data" / "manifests"
    manifest_dir.mkdir(parents=True, exist_ok=True)
    (manifest_dir / "AAA.json").write_text('{"symbol":"AAA","hash":"raw"}', encoding="utf-8")
    if db_marker:
        state_dir = root / "state"
        state_dir.mkdir(parents=True, exist_ok=True)
        (state_dir / "engine.sqlite3").write_bytes(b"final-holdout-marker")

    dates = pd.bdate_range("2026-01-01", periods=90)
    symbols = ("AAA", "BBB", "CCC", "TZA")
    rows: list[dict[str, object]] = []
    for date_index, date in enumerate(dates):
        for symbol_index, symbol in enumerate(symbols):
            phase = date_index / 6.0 + symbol_index * 0.7
            role = "leveraged_inverse_etf" if symbol == "TZA" else "stock"
            return_5 = 0.025 * math.sin(phase)
            return_20 = 0.030 * math.cos(phase / 1.7)
            close_position = 0.50 + 0.40 * math.sin(phase / 2.1)
            atr = 0.018 + 0.012 * abs(math.sin(phase / 1.9))
            score = 0.70 * return_5 + 0.05 * (close_position - 0.50) - 0.25 * atr
            bull_forward = score + 0.010 * (1 if (date_index + symbol_index) % 3 == 0 else -0.4)
            bear_forward = -bull_forward + 0.004 * (
                1 if (date_index + symbol_index) % 4 == 0 else -0.2
            )
            bull_forward_20 = return_20 + 0.008 * (
                1 if (date_index + symbol_index) % 4 in {0, 1} else -0.6
            )
            bear_forward_20 = -bull_forward_20 + 0.003 * (
                1 if (date_index + symbol_index) % 5 == 0 else -0.3
            )
            rows.append(
                {
                    "Date": date,
                    "symbol": symbol,
                    "role": role,
                    "Open": 100.0 + date_index + symbol_index,
                    "High": 101.0 + date_index + symbol_index,
                    "Low": 99.0 + date_index + symbol_index,
                    "Close": 100.5 + date_index + symbol_index,
                    "Volume": 1_000_000 + date_index * 1000 + symbol_index,
                    "return_5": return_5,
                    "return_20": return_20,
                    "momentum_20_percentile_252": ((date_index + symbol_index) % 20) / 19,
                    "trend_persistence_20": return_20 * 2.0,
                    "range_position_20": close_position,
                    "distance_prior_high_20": 0.02 - return_5,
                    "atr_pct_14": atr,
                    "realized_vol_20": atr * 1.2,
                    "volatility_expansion_20_63": 0.9 + atr * 8.0,
                    "close_position": close_position,
                    "upper_wick_pct": 0.15 + 0.03 * symbol_index,
                    "lower_wick_pct": 0.20 - 0.02 * symbol_index,
                    "relative_volume_20": 1.0 + 0.01 * date_index,
                    "dollar_volume": 8_000_000 + 50_000 * date_index,
                    "relative_return_vs_spy_20": return_20 - 0.005,
                    "rolling_corr_vs_spy_63": 0.25,
                    "relative_return_vs_sector_20": return_20 - 0.002 * symbol_index,
                    "sector_momentum_rank_20": ((date_index + 2 * symbol_index) % 23) / 22,
                    "breadth_advance_pct": 0.45 + 0.01 * ((date_index + symbol_index) % 5),
                    "breadth_dispersion_20": 0.08 + atr,
                    "inverse_confirmation_iwm_tza_63": 0.80 if symbol == "TZA" else 0.10,
                    "relationship_mutual_info_iwm_tza_63": 1.4 if symbol == "TZA" else 0.2,
                    "relationship_breakdown_iwm_tza_63": 0.0,
                    "market_regime_trend_score": 0.3 + return_20,
                    "market_regime_volatility_score": atr * 10.0,
                    "rsi_14": 50.0 + return_5 * 100.0,
                    "label_bull_forward_return_5": bull_forward,
                    "label_bull_positive_return_5": int(bull_forward > 0.0),
                    "label_bull_mfe_5": max(bull_forward + 0.030, 0.002),
                    "label_bull_mae_5": min(bull_forward - 0.030, -0.002),
                    "label_bull_target_before_stop_5": int(bull_forward > -0.004),
                    "label_bull_time_to_target_5": 3 if bull_forward > -0.004 else 6,
                    "label_bull_time_to_stop_5": 6 if bull_forward > -0.004 else 3,
                    "label_bear_forward_return_5": bear_forward,
                    "label_bear_positive_return_5": int(bear_forward > 0.0),
                    "label_bear_mfe_5": max(bear_forward + 0.030, 0.002),
                    "label_bear_mae_5": min(bear_forward - 0.030, -0.002),
                    "label_bear_target_before_stop_5": int(bear_forward > -0.004),
                    "label_bear_time_to_target_5": 3 if bear_forward > -0.004 else 6,
                    "label_bear_time_to_stop_5": 6 if bear_forward > -0.004 else 3,
                    "label_end_date_5": dates[min(date_index + 5, len(dates) - 1)],
                    "label_bull_forward_return_20": bull_forward_20,
                    "label_bull_positive_return_20": int(bull_forward_20 > 0.0),
                    "label_bull_mfe_20": max(bull_forward_20 + 0.045, 0.004),
                    "label_bull_mae_20": min(bull_forward_20 - 0.045, -0.004),
                    "label_bull_target_before_stop_20": int(bull_forward_20 > -0.006),
                    "label_bull_time_to_target_20": 8 if bull_forward_20 > -0.006 else 22,
                    "label_bull_time_to_stop_20": 22 if bull_forward_20 > -0.006 else 5,
                    "label_bear_forward_return_20": bear_forward_20,
                    "label_bear_positive_return_20": int(bear_forward_20 > 0.0),
                    "label_bear_mfe_20": max(bear_forward_20 + 0.045, 0.004),
                    "label_bear_mae_20": min(bear_forward_20 - 0.045, -0.004),
                    "label_bear_target_before_stop_20": int(bear_forward_20 > -0.006),
                    "label_bear_time_to_target_20": 8 if bear_forward_20 > -0.006 else 22,
                    "label_bear_time_to_stop_20": 22 if bear_forward_20 > -0.006 else 5,
                    "label_end_date_20": dates[min(date_index + 20, len(dates) - 1)],
                }
            )
    modeling = pd.DataFrame(rows)
    feature_columns = [column for column in modeling.columns if not column.startswith("label_")]
    features = modeling[feature_columns].copy()
    features.to_parquet(feature_dir / "universehash_testfeatures_features.parquet", index=False)
    modeling.to_parquet(feature_dir / "universehash_testfeatures_modeling.parquet", index=False)


def _write_policy_candidate_feature_data(root: Path) -> None:
    feature_dir = root / "data" / "features"
    feature_dir.mkdir(parents=True, exist_ok=True)
    manifest_dir = root / "data" / "manifests"
    manifest_dir.mkdir(parents=True, exist_ok=True)
    (manifest_dir / "AAA.json").write_text('{"symbol":"AAA","hash":"raw"}', encoding="utf-8")

    dates = pd.bdate_range("2025-01-02", periods=190)
    symbols = ("AAA", "BBB", "CCC", "TZA")
    feature_rows: list[pd.DataFrame] = []
    label_rows: list[pd.DataFrame] = []
    for symbol_index, symbol in enumerate(symbols):
        role = "leveraged_inverse_etf" if symbol == "TZA" else "stock"
        rows: list[dict[str, object]] = []
        base = 80.0 + symbol_index * 9.0
        for date_index, date in enumerate(dates):
            phase = date_index / 4.5 + symbol_index * 0.8
            close = base + date_index * 0.04 + 4.0 * math.sin(phase)
            open_price = close + 0.35 * math.sin(phase * 1.7)
            high = max(open_price, close) + 1.15 + 0.75 * abs(math.sin(phase * 0.9))
            low = min(open_price, close) - 1.15 - 0.75 * abs(math.cos(phase * 0.8))
            return_5 = 0.025 * math.sin(phase)
            return_20 = 0.035 * math.sin(phase / 1.8)
            rows.append(
                {
                    "Date": date,
                    "symbol": symbol,
                    "role": role,
                    "sector": "test",
                    "Open": open_price,
                    "High": high,
                    "Low": low,
                    "Close": close,
                    "Volume": 1_500_000 + date_index * 2000 + symbol_index * 100,
                    "return_5": return_5,
                    "return_20": return_20,
                    "momentum_20_percentile_252": ((date_index + symbol_index) % 30) / 29,
                    "trend_persistence_20": return_20 * 1.8,
                    "range_position_20": 0.50 + 0.35 * math.sin(phase / 2.0),
                    "distance_prior_high_20": 0.03 - return_5,
                    "atr_pct_14": 0.020 + 0.010 * abs(math.sin(phase / 2.0)),
                    "realized_vol_20": 0.025 + 0.010 * abs(math.cos(phase / 2.5)),
                    "volatility_expansion_20_63": 0.85 + 0.10 * abs(math.sin(phase)),
                    "close_position": 0.50 + 0.30 * math.sin(phase / 1.5),
                    "upper_wick_pct": 0.12 + 0.03 * abs(math.sin(phase)),
                    "lower_wick_pct": 0.12 + 0.03 * abs(math.cos(phase)),
                    "relative_volume_20": 1.0 + 0.05 * math.sin(phase / 3.0),
                    "dollar_volume": 10_000_000 + date_index * 75_000,
                    "relative_return_vs_spy_20": return_20 - 0.003,
                    "rolling_corr_vs_spy_63": 0.30 + 0.05 * math.sin(phase / 3.0),
                    "relative_return_vs_sector_20": return_20 - 0.002 * symbol_index,
                    "sector_momentum_rank_20": ((date_index + symbol_index) % 23) / 22,
                    "breadth_advance_pct": 0.45 + 0.02 * math.sin(phase / 2.0),
                    "breadth_dispersion_20": 0.08 + 0.01 * abs(math.cos(phase)),
                    "market_regime_trend_score": 0.35 + return_20,
                    "market_regime_volatility_score": 0.20 + abs(return_5),
                    "rsi_14": 50.0 + return_5 * 100.0,
                }
            )
        feature_rows.append(pd.DataFrame(rows))

    features = (
        pd.concat(feature_rows, ignore_index=True)
        .sort_values(["Date", "symbol"])
        .reset_index(drop=True)
    )
    for symbol, group in features.groupby("symbol", sort=True):
        labels = build_symbol_labels(
            group.set_index("Date"),
            LabelConfig(horizons=(20,), target_atr_multiple=2.0, stop_atr_multiple=1.0),
        )
        label_rows.append(
            labels.assign(symbol=symbol).reset_index().sort_values(["Date", "symbol"])
        )
    labels = pd.concat(label_rows, ignore_index=True)
    modeling = features.merge(labels, on=["Date", "symbol"], how="inner", validate="one_to_one")
    features.to_parquet(feature_dir / "universehash_policyfeatures_features.parquet", index=False)
    modeling.to_parquet(feature_dir / "universehash_policyfeatures_modeling.parquet", index=False)


def _write_policy_candidate_config(root: Path) -> Path:
    config_dir = root / "configs" / "signal_discovery"
    config_dir.mkdir(parents=True, exist_ok=True)
    path = config_dir / "policy_candidate_v1.yaml"
    path.write_text(
        f"""
schema_version: multi_angle_signal_discovery_v1
enabled_hypotheses:
  - sector_rotation_buy_20d
  - {SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_HYPOTHESIS_ID}
  - {SECTOR_ROTATION_BUY_ORDINARY_TIME_EXIT_HYPOTHESIS_ID}
horizons: [20]
product_scopes: [POOLED, ORDINARY, LEVERAGED_INVERSE]
model_families:
  - hist_gradient_boosting
minimum_training_samples: 80
minimum_calibration_samples: 30
minimum_holdout_samples: 30
max_selected_features: 12
analog_count: 3
candidate_cap_per_hypothesis: 8
random_seed: 23
research_start: "2025-01-02"
research_end:
costs:
  round_trip_bps: 5.0
selection_policy:
  reference: signal_discovery_policy_test_v1
  probability: 0.10
  target_before_stop: 0.10
  expected_return: -0.10
  signal_score: 0.05
ood_policy:
  reference: prediction_ood_governance_test_v1
  feature_rate_limit: 1.00
""",
        encoding="utf-8",
    )
    return path


def _write_blocked_analog_fixture(root: Path) -> None:
    feature_dir = root / "data" / "features"
    feature_dir.mkdir(parents=True, exist_ok=True)
    generation_id = "signal_discovery_20260628T193012+0000_fixture"
    generation_dir = root / "artifacts" / "signal_discovery" / generation_id
    generation_dir.mkdir(parents=True, exist_ok=True)
    labels = {
        "directional_return": "label_bull_forward_return_20",
        "positive_return": "label_bull_positive_return_20",
        "mfe": "label_bull_mfe_20",
        "mae": "label_bull_mae_20",
        "target_before_stop": "label_bull_target_before_stop_20",
        "time_to_target": "label_bull_time_to_target_20",
        "time_to_stop": "label_bull_time_to_stop_20",
        "label_end_date": "label_end_date_20",
    }
    selected_features = ["feature_a", "feature_b", "label_bull_forward_return_20"]
    target_payload = json.dumps({"feature_a": 0.0, "feature_b": 0.0})
    candidates = pd.DataFrame(
        [
            {
                "signal_id": "target-tza",
                "generation_id": generation_id,
                "as_of_date": "2026-06-26",
                "ticker": "TZA",
                "symbol": "TZA",
                "direction": "Bullish",
                "action": "NO SIGNAL",
                "decision": "NO_SIGNAL",
                "candidate_status": "RESEARCH_ONLY",
                "archetype": "Sector Rotation",
                "archetype_id": "sector_rotation",
                "hypothesis_id": "sector_rotation_buy_20d",
                "model_id": "sector_rotation_buy_20d:extra_trees",
                "model_family": "extra_trees",
                "scope": "LEVERAGED_INVERSE",
                "product_class_scope": "LEVERAGED_INVERSE",
                "selected_feature_values_json": target_payload,
                "signal_score": 0.90,
                "no_signal_reason": "target_before_stop_probability_below_threshold",
                "rejection_reason": "",
            },
            {
                "signal_id": "target-soxs",
                "generation_id": generation_id,
                "as_of_date": "2026-06-26",
                "ticker": "SOXS",
                "symbol": "SOXS",
                "direction": "Bullish",
                "action": "BUY",
                "decision": "REJECTED_BY_OOD",
                "candidate_status": "REJECTED_BY_OOD",
                "archetype": "Sector Rotation",
                "archetype_id": "sector_rotation",
                "hypothesis_id": "sector_rotation_buy_20d",
                "model_id": "sector_rotation_buy_20d:extra_trees",
                "model_family": "extra_trees",
                "scope": "LEVERAGED_INVERSE",
                "product_class_scope": "LEVERAGED_INVERSE",
                "selected_feature_values_json": json.dumps({"feature_a": 0.02, "feature_b": 0.01}),
                "signal_score": 0.80,
                "no_signal_reason": "",
                "rejection_reason": "ood_feature_rate_above_limit",
            },
            {
                "signal_id": "target-prob",
                "generation_id": generation_id,
                "as_of_date": "2026-06-26",
                "ticker": "AAA",
                "symbol": "AAA",
                "direction": "Bullish",
                "action": "NO SIGNAL",
                "decision": "NO_SIGNAL",
                "candidate_status": "RESEARCH_ONLY",
                "archetype": "Sector Rotation",
                "archetype_id": "sector_rotation",
                "hypothesis_id": "sector_rotation_buy_20d",
                "model_id": "sector_rotation_buy_20d:extra_trees",
                "model_family": "extra_trees",
                "scope": "ORDINARY",
                "product_class_scope": "ORDINARY",
                "selected_feature_values_json": json.dumps({"feature_a": 0.03, "feature_b": 0.02}),
                "signal_score": 0.70,
                "no_signal_reason": "probability_below_threshold",
                "rejection_reason": "",
            },
        ]
    )
    hypotheses = pd.DataFrame(
        [
            {
                "generation_id": generation_id,
                "hypothesis_id": "sector_rotation_buy_20d",
                "archetype": "Sector Rotation",
                "direction": "BUY",
                "horizon": 20,
                "family": "extra_trees",
                "status": "CANDIDATE",
                "outcome_labels": json.dumps(labels, sort_keys=True),
                "selected_features": json.dumps(selected_features),
            }
        ]
    )
    modeling_rows = [
        ("2026-06-22", "TZA", "leveraged_inverse_etf", 0.01, 0.00, 0.04, 0.09, -0.02, 1),
        ("2026-06-19", "SQQQ", "leveraged_inverse_etf", 0.30, 0.25, -0.02, 0.03, -0.05, 0),
        ("2026-06-18", "AAA", "stock", 0.00, 0.00, 0.03, 0.06, -0.01, 1),
        ("2026-06-17", "BBB", "stock", 0.04, 0.02, -0.01, 0.02, -0.04, 0),
        ("2026-06-26", "TZA", "leveraged_inverse_etf", 0.00, 0.00, 0.99, 1.20, -0.90, 1),
        ("2026-06-29", "SOXS", "leveraged_inverse_etf", 0.02, 0.01, 0.88, 1.10, -0.80, 1),
    ]
    modeling = pd.DataFrame(
        [
            {
                "Date": date,
                "symbol": symbol,
                "role": role,
                "feature_a": feature_a,
                "feature_b": feature_b,
                "market_regime_cluster_expanding": "fixture",
                "label_bull_forward_return_20": forward_return,
                "label_bull_positive_return_20": int(forward_return > 0.0),
                "label_bull_mfe_20": mfe,
                "label_bull_mae_20": mae,
                "label_bull_target_before_stop_20": tbs,
                "label_bull_time_to_target_20": 5,
                "label_bull_time_to_stop_20": 8,
                "label_end_date_20": "2026-07-24",
            }
            for (
                date,
                symbol,
                role,
                feature_a,
                feature_b,
                forward_return,
                mfe,
                mae,
                tbs,
            ) in modeling_rows
        ]
    )
    pd.DataFrame([{"generation_id": generation_id}]).to_csv(
        generation_dir / "summary.csv", index=False
    )
    hypotheses.to_csv(generation_dir / "hypotheses.csv", index=False)
    candidates.to_csv(generation_dir / "candidates.csv", index=False)
    candidates.loc[candidates["decision"].eq("NO_SIGNAL")].to_csv(
        generation_dir / "no_signal.csv", index=False
    )
    candidates.loc[candidates["decision"].str.startswith("REJECTED")].to_csv(
        generation_dir / "rejected.csv", index=False
    )
    pd.DataFrame().to_csv(generation_dir / "selected_candidates.csv", index=False)
    pd.DataFrame().to_csv(generation_dir / "footprint_evidence.csv", index=False)
    pd.DataFrame().to_csv(generation_dir / "historical_analogs.csv", index=False)
    pd.DataFrame().to_csv(generation_dir / "score_components.csv", index=False)
    pd.DataFrame().to_csv(generation_dir / "gate_results.csv", index=False)
    (generation_dir / "metadata.json").write_text(
        json.dumps({"generation_id": generation_id}, sort_keys=True), encoding="utf-8"
    )
    (generation_dir.parent / "latest.json").write_text(
        json.dumps({"generation_id": generation_id}, sort_keys=True), encoding="utf-8"
    )
    modeling.to_parquet(feature_dir / "fixturehash_blockedanalog_modeling.parquet", index=False)


def _analog_row(
    rank: int,
    *,
    target: str = "target-robust",
    ticker: str,
    year: int,
    day: int,
    regime: str,
    forward_return: float,
    tbs: bool,
    mae: float = -0.03,
    same_symbol: bool = False,
) -> dict[str, object]:
    return {
        "target_row_id": target,
        "target_signal_id": target,
        "generation_id": "robustness-fixture",
        "target_as_of_date": "2026-06-26",
        "target_ticker": "SOXS",
        "target_direction": "Bearish",
        "target_archetype": "Breakout / Breakdown",
        "target_action": "NO SIGNAL",
        "target_status": "RESEARCH_ONLY",
        "target_score": 0.70,
        "target_blocker_reason": "target_before_stop_probability_below_threshold",
        "target_hypothesis_id": "breakdown_sell_10d",
        "target_model_id": "breakdown_sell_10d:extra_trees",
        "target_product_scope": "LEVERAGED_INVERSE",
        "target_selection_reason": "top_bearish",
        "analog_rank": rank,
        "analog_date": f"{year}-01-{day:02d}",
        "analog_ticker": ticker,
        "analog_scope": "LEVERAGED_INVERSE",
        "analog_direction": "Bearish",
        "analog_archetype": "Breakout / Breakdown",
        "analog_pool": "same_scope",
        "similarity_score": max(0.05, 0.90 - rank * 0.01),
        "distance_score": float(rank),
        "same_symbol": same_symbol,
        "same_product_scope": True,
        "same_archetype": True,
        "same_direction": True,
        "market_regime": regime,
        "forward_return": forward_return,
        "MFE": max(forward_return + 0.05, 0.01),
        "MAE": mae,
        "target_before_stop_result": "target before stop" if tbs else "stop before target",
        "analog_would_have_passed_current_thresholds": "not_available_existing_artifacts_only",
        "analog_rejection_reason": "not_available_existing_artifacts_only",
        "outcome_labels_used_for_explanation_only": True,
    }


def _concentrated_decay_analogs() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for rank in range(1, 11):
        rows.append(
            _analog_row(
                rank,
                ticker="SOXS",
                year=2026,
                day=rank,
                regime="event_cluster",
                forward_return=0.12,
                tbs=True,
                same_symbol=True,
            )
        )
    for rank in range(11, 51):
        rows.append(
            _analog_row(
                rank,
                ticker="SOXS" if rank <= 48 else "TZA",
                year=2025 if rank <= 30 else 2024,
                day=((rank - 1) % 28) + 1,
                regime="other_regime",
                forward_return=-0.08 if rank % 2 else 0.02,
                tbs=False,
                mae=-0.25 if rank == 25 else -0.04,
                same_symbol=rank <= 48,
            )
        )
    return pd.DataFrame(rows)


def _diversified_supportive_analogs() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for rank in range(1, 51):
        rows.append(
            _analog_row(
                rank,
                ticker=f"SYM{rank % 12}",
                year=2017 + (rank % 8),
                day=((rank - 1) % 28) + 1,
                regime=f"regime_{rank % 4}",
                forward_return=0.04 + (rank % 3) * 0.002,
                tbs=rank % 4 != 0,
                mae=-0.025,
                same_symbol=False,
            )
        )
    return pd.DataFrame(rows)


def _diversified_decay_analogs() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for rank in range(1, 11):
        rows.append(
            _analog_row(
                rank,
                ticker=f"SYM{rank}",
                year=2016 + rank,
                day=rank,
                regime=f"regime_{rank % 5}",
                forward_return=0.08,
                tbs=True,
                same_symbol=False,
            )
        )
    for rank in range(11, 51):
        rows.append(
            _analog_row(
                rank,
                ticker=f"SYM{rank % 20}",
                year=2017 + (rank % 8),
                day=((rank - 1) % 28) + 1,
                regime=f"regime_{rank % 5}",
                forward_return=-0.05,
                tbs=False,
                same_symbol=False,
            )
        )
    return pd.DataFrame(rows)


@pytest.fixture()
def signal_discovery_root(tmp_path: Path) -> tuple[Path, Path]:
    root = tmp_path / "dev"
    _write_universe(root)
    _write_feature_data(root, db_marker=True)
    config = _write_config(root)
    return root, config


def test_hypothesis_registry_loads_required_archetypes_without_privileged_rsi() -> None:
    archetypes = default_archetype_registry()
    hypotheses = default_hypothesis_registry()

    assert len(archetypes) == 10
    assert "risk_on_risk_off" in archetypes
    assert "failed_move_liquidity_trap" in archetypes
    assert any(spec.direction == "BUY" for spec in hypotheses.values())
    assert any(spec.direction == "SELL_SHORT" for spec in hypotheses.values())
    for spec in hypotheses.values():
        assert spec.archetype_id
        assert spec.direction in {"BUY", "SELL_SHORT"}
        assert spec.horizon in {3, 5, 10, 20}
        assert spec.outcome_labels
        assert spec.candidate_feature_families
        assert "rsi_family" not in spec.required_feature_families
        assert spec.validation_policy["feature_screen_split"] == "training_only"
        assert spec.validation_policy["calibration_split"] == "calibration_only"
        assert spec.validation_policy["holdout_usage"] == "evaluation_only"

    candidate = hypotheses[SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_HYPOTHESIS_ID]
    assert candidate.eligible_product_scopes == ("ORDINARY",)
    assert candidate.direction == "BUY"
    assert candidate.target_stop_policy_status == "EXPERIMENTAL_CANDIDATE"
    assert candidate.target_stop_policy_id == SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    baseline = hypotheses["sector_rotation_buy_20d"]
    assert baseline.target_stop_policy_status == "DEFAULT_BASELINE"
    assert baseline.target_stop_policy_target_multiple == 2.0
    assert baseline.target_stop_policy_stop_multiple == 1.0
    time_exit = hypotheses[SECTOR_ROTATION_BUY_ORDINARY_TIME_EXIT_HYPOTHESIS_ID]
    assert time_exit.eligible_product_scopes == ("ORDINARY",)
    assert time_exit.direction == "BUY"
    assert time_exit.archetype_id == "sector_rotation"
    assert time_exit.time_exit_label_schema_version == TIME_EXIT_UTILITY_LABEL_SCHEMA_VERSION
    assert "time_exit_utility" in time_exit.outcome_labels
    assert "expected_time_exit_utility" in time_exit.model_tasks


def test_target_stop_policy_candidate_selection_uses_calibration_only_rule() -> None:
    selection = select_sector_rotation_buy_ordinary_policy_candidate()
    assert selection.status == "CALIBRATION_SUPPORTED_POLICY_CANDIDATE"
    assert selection.selected is not None
    assert selection.selected.target_multiple == 2.0
    assert selection.selected.stop_multiple == 1.25
    assert selection.selected_policy is not None
    assert selection.selected_policy.governance_status == "EXPERIMENTAL_CANDIDATE"
    assert selection.selected_policy.policy_id == SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID

    registry = {policy.policy_id: policy for policy in target_stop_policy_registry()}
    assert registry[SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID].stop_multiple == 1.25
    assert registry[SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID].target_multiple == 2.0

    baseline = SECTOR_ROTATION_BUY_ORDINARY_CALIBRATION_EVIDENCE[0]
    unsupported = (
        baseline,
        CalibrationPolicyEvidence(
            horizon=20,
            target_multiple=2.0,
            stop_multiple=1.25,
            row_count=baseline.row_count,
            target_before_stop_hit_rate=baseline.target_before_stop_hit_rate,
            stop_before_target_rate=baseline.stop_before_target_rate + 0.01,
            unresolved_rate=baseline.unresolved_rate,
            average_forward_return=baseline.average_forward_return,
            median_forward_return=baseline.median_forward_return,
            average_mfe=baseline.average_mfe,
            average_mae=baseline.average_mae,
            worst_mae=baseline.worst_mae,
            expected_r=baseline.expected_r,
            cost_adjusted_utility=baseline.cost_adjusted_utility,
            symbol_concentration=baseline.symbol_concentration,
            year_concentration=baseline.year_concentration,
            regime_concentration=baseline.regime_concentration,
        ),
    )
    no_candidate = select_sector_rotation_buy_ordinary_policy_candidate(unsupported)
    assert no_candidate.status == "NO_CALIBRATION_SUPPORTED_POLICY_CANDIDATE"
    assert no_candidate.selected_policy is None


def test_candidate_policy_derived_labels_are_separate_from_baseline_labels(tmp_path: Path) -> None:
    root = tmp_path / "policy-labels"
    _write_universe(root)
    _write_policy_candidate_feature_data(root)
    modeling_path = next((root / "data" / "features").glob("*_modeling.parquet"))
    modeling = pd.read_parquet(modeling_path)
    before_baseline = modeling["label_bull_target_before_stop_20"].copy()

    candidate_policy = next(
        policy
        for policy in target_stop_policy_registry()
        if policy.policy_id == SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    )
    augmented, derived = augment_model_frame_with_policy_outcomes(modeling, (candidate_policy,))
    label_columns = candidate_policy_outcome_labels(candidate_policy)

    assert label_columns["target_before_stop"] in augmented.columns
    assert label_columns["time_to_target"] in augmented.columns
    assert not derived.empty
    assert set(derived["target_stop_policy_id"]) == {
        SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    }
    pd.testing.assert_series_equal(
        modeling["label_bull_target_before_stop_20"],
        before_baseline,
        check_names=False,
    )
    assert label_columns["target_before_stop"] != "label_bull_target_before_stop_20"
    assert augmented[label_columns["target_before_stop"]].dropna().isin([0.0, 1.0]).all()


def test_time_exit_utility_labels_are_policy_separate_and_label_side(
    tmp_path: Path,
) -> None:
    root = tmp_path / "time-exit-labels"
    _write_universe(root)
    _write_policy_candidate_feature_data(root)
    modeling_path = next((root / "data" / "features").glob("*_modeling.parquet"))
    modeling = pd.read_parquet(modeling_path)
    before_modeling = modeling_path.read_bytes()
    policies = target_stop_policy_registry()

    augmented, labels = augment_model_frame_with_time_exit_utility_labels(
        modeling,
        policies,
        cost_return=0.0005,
    )

    assert not labels.empty
    assert set(labels["scope"].astype(str)) == {"ORDINARY"}
    assert "TZA" not in set(labels["symbol"].astype(str))
    assert TIME_EXIT_BASELINE_POLICY_ALIAS in set(labels["target_stop_policy_alias"].astype(str))
    assert SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID in set(
        labels["target_stop_policy_id"].astype(str)
    )
    assert set(labels["schema_version"].astype(str)) == {TIME_EXIT_UTILITY_LABEL_SCHEMA_VERSION}
    assert labels["time_exit_positive_after_cost_20d"].dropna().isin([0.0, 1.0]).all()
    assert labels["profitable_despite_failed_tbs_20d"].dropna().isin([0.0, 1.0]).all()
    assert labels["early_adverse_recovery_20d"].dropna().isin([0.0, 1.0]).all()

    candidate_policy = next(
        policy
        for policy in policies
        if policy.policy_id == SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    )
    label_map = time_exit_utility_outcome_labels(candidate_policy)
    for key in (
        "positive_return",
        "directional_return",
        "time_exit_utility",
        "profitable_despite_failed_tbs",
        "early_adverse_recovery",
    ):
        assert label_map[key] in augmented.columns

    sample = (
        labels.loc[
            labels["symbol"].astype(str).eq("AAA")
            & labels["target_stop_policy_id"].astype(str).eq(candidate_policy.policy_id)
            & labels["entry_price"].notna()
        ]
        .sort_values("Date")
        .iloc[0]
    )
    symbol_frame = (
        modeling.loc[modeling["symbol"].astype(str).eq("AAA")]
        .sort_values("Date")
        .reset_index(drop=True)
    )
    position = int(
        symbol_frame.index[pd.to_datetime(symbol_frame["Date"]).eq(pd.Timestamp(sample["Date"]))][0]
    )
    assert sample["entry_price"] == pytest.approx(symbol_frame.iloc[position + 1]["Open"])
    assert sample["exit_price"] == pytest.approx(symbol_frame.iloc[position + 20]["Close"])
    assert sample["time_exit_net_return_20d"] == pytest.approx(
        sample["time_exit_gross_return_20d"] - 0.0005
    )
    assert bool(sample["profitable_despite_failed_tbs_20d"]) == (
        bool(sample["target_before_stop"] < 0.5) and bool(sample["time_exit_net_return_20d"] > 0.0)
    )
    assert modeling_path.read_bytes() == before_modeling


def test_time_exit_utility_guard_blocks_label_side_columns() -> None:
    frame = pd.DataFrame(
        {
            "Date": pd.date_range("2026-01-01", periods=2),
            "symbol": ["AAA", "AAA"],
            "return_20": [0.01, 0.02],
            "time_exit_net_return_20d": [0.01, -0.01],
            "time_exit_utility_20d": [0.5, -0.2],
            "profitable_despite_failed_tbs_20d": [1.0, 0.0],
            "early_adverse_recovery_20d": [0.0, 1.0],
            "label_custom_time_exit": [1.0, 0.0],
        }
    )

    assert numeric_feature_columns(frame) == ["return_20"]
    with pytest.raises(ValueError, match="Label-side columns"):
        reject_label_columns(["return_20", "time_exit_utility_20d"])


def test_time_exit_utility_handles_near_zero_mae_safely() -> None:
    dates = pd.bdate_range("2026-01-02", periods=26)
    frame = pd.DataFrame(
        {
            "Date": dates,
            "symbol": "AAA",
            "role": "stock",
            "sector": "test",
            "Open": 100.0,
            "High": 100.0,
            "Low": 100.0,
            "Close": 100.0,
            "Volume": 1_000_000,
            "atr_pct_14": 0.0,
            "label_end_date_20": dates.to_series(index=range(len(dates))).shift(-20),
        }
    )

    labels = build_time_exit_utility_labels(
        frame,
        target_stop_policy_registry(),
        cost_return=0.0005,
    )
    finite = pd.to_numeric(labels["time_exit_utility_20d"], errors="coerce").dropna()
    assert not finite.empty
    assert finite.map(math.isfinite).all()
    assert time_exit_quality_bucket(0.025) == "STRONG_POSITIVE_TIME_EXIT"
    assert time_exit_quality_bucket(0.001) == "MODEST_POSITIVE_TIME_EXIT"
    assert time_exit_quality_bucket(-0.001) == "FLAT_TIME_EXIT"
    assert time_exit_quality_bucket(-0.01) == "NEGATIVE_TIME_EXIT"
    assert time_exit_quality_bucket(-0.06) == "SEVERE_NEGATIVE_TIME_EXIT"


def test_time_exit_calibration_summary_separates_calibration_and_holdout(
    tmp_path: Path,
) -> None:
    root = tmp_path / "time-exit-summary"
    _write_universe(root)
    _write_policy_candidate_feature_data(root)
    modeling = pd.read_parquet(next((root / "data" / "features").glob("*_modeling.parquet")))
    labels = build_time_exit_utility_labels(
        modeling,
        target_stop_policy_registry(),
        cost_return=0.0005,
    )

    summary = time_exit_utility_calibration_summary_frame(labels)

    assert not summary.empty
    assert {"calibration_only", "DEVELOPMENT_HOLDOUT_DIAGNOSTIC_ONLY"}.issubset(
        set(summary["evidence_split"].astype(str))
    )
    calibration = summary.loc[summary["evidence_split"].astype(str).eq("calibration_only")]
    holdout = summary.loc[
        summary["evidence_split"].astype(str).eq("DEVELOPMENT_HOLDOUT_DIAGNOSTIC_ONLY")
    ]
    assert not calibration["development_holdout_diagnostic_only"].astype(bool).any()
    assert holdout["development_holdout_diagnostic_only"].astype(bool).all()
    assert calibration["sample_count"].min() > 0
    assert holdout["sample_count"].min() > 0


def test_signal_discovery_persists_target_stop_policy_candidate_artifacts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "policy-discovery"
    _write_universe(root)
    _write_policy_candidate_feature_data(root)
    config = _write_policy_candidate_config(root)
    model_artifact = root / "artifacts" / "models" / "model.joblib"
    model_artifact.parent.mkdir(parents=True, exist_ok=True)
    model_artifact.write_text("model-artifact-do-not-touch", encoding="utf-8")
    before_model_artifact = model_artifact.read_text(encoding="utf-8")
    modeling_path = next((root / "data" / "features").glob("*_modeling.parquet"))
    before_modeling = modeling_path.read_bytes()

    def fail_download(*_: object, **__: object) -> None:
        raise AssertionError("Policy candidate discovery attempted an FMP request")

    monkeypatch.setattr("swing_rsi.data.loader.download_daily", fail_download)
    run_signal_discovery(root, config_path=config)
    frames = load_signal_discovery_frames(root)

    assert not frames["target_stop_policy_registry"].empty
    assert not frames["calibration_selection"].empty
    assert not frames["sector_rotation_buy_ordinary_policy_comparison"].empty
    assert not frames["derived_policy_outcomes"].empty
    assert not frames["signal_discovery_policy_comparison"].empty
    assert not frames["time_exit_utility_labels"].empty
    assert not frames["time_exit_utility_calibration_summary"].empty
    assert not frames["time_exit_utility_signal_rows"].empty
    assert not frames["time_exit_utility_policy_comparison"].empty
    assert (root / "artifacts" / "signal_discovery" / "latest.json").exists()
    registry_ids = set(frames["target_stop_policy_registry"]["policy_id"].astype(str))
    assert SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID in registry_ids
    hypotheses = frames["hypotheses"]
    candidate_hypothesis = hypotheses.loc[
        hypotheses["hypothesis_id"]
        .astype(str)
        .eq(SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_HYPOTHESIS_ID)
    ]
    assert not candidate_hypothesis.empty
    assert set(candidate_hypothesis["target_stop_policy_id"].astype(str)) == {
        SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    }
    assert set(candidate_hypothesis["target_stop_policy_status"].astype(str)) == {
        "EXPERIMENTAL_CANDIDATE"
    }
    candidate_rows = frames["candidates"].loc[
        frames["candidates"]["hypothesis_id"]
        .astype(str)
        .eq(SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_HYPOTHESIS_ID)
    ]
    assert not candidate_rows.empty
    assert set(candidate_rows["product_class_scope"].astype(str)) == {"ORDINARY"}
    assert set(candidate_rows["target_stop_policy_id"].astype(str)) == {
        SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    }
    assert "LIVE_ACTIONABLE" not in set(candidate_rows["candidate_status"].astype(str))
    gate_rows = frames["gate_results"].loc[
        frames["gate_results"]["hypothesis_id"]
        .astype(str)
        .eq(SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_HYPOTHESIS_ID)
    ]
    assert not gate_rows.empty
    assert set(gate_rows["target_stop_policy_id"].astype(str)) == {
        SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    }
    comparison = frames["signal_discovery_policy_comparison"]
    assert SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID in set(
        comparison["target_stop_policy_id"].astype(str)
    )
    time_exit_hypothesis = hypotheses.loc[
        hypotheses["hypothesis_id"]
        .astype(str)
        .eq(SECTOR_ROTATION_BUY_ORDINARY_TIME_EXIT_HYPOTHESIS_ID)
    ]
    assert not time_exit_hypothesis.empty
    assert set(time_exit_hypothesis["target_stop_policy_id"].astype(str)) == {
        SECTOR_ROTATION_BUY_ORDINARY_CANDIDATE_POLICY_ID
    }
    assert set(time_exit_hypothesis["time_exit_label_schema_version"].astype(str)) == {
        TIME_EXIT_UTILITY_LABEL_SCHEMA_VERSION
    }
    time_exit_rows = frames["candidates"].loc[
        frames["candidates"]["hypothesis_id"]
        .astype(str)
        .eq(SECTOR_ROTATION_BUY_ORDINARY_TIME_EXIT_HYPOTHESIS_ID)
    ]
    assert not time_exit_rows.empty
    assert set(time_exit_rows["product_class_scope"].astype(str)) == {"ORDINARY"}
    assert "LIVE_ACTIONABLE" not in set(time_exit_rows["candidate_status"].astype(str))
    assert (
        time_exit_rows["not_live_actionable_reason"]
        .astype(str)
        .str.contains(
            "Time-exit utility is diagnostic",
            regex=False,
        )
        .all()
    )
    for column in (
        "time_exit_positive_probability",
        "expected_time_exit_return",
        "expected_time_exit_utility",
        "profitable_despite_failed_tbs_probability",
        "early_adverse_recovery_probability",
    ):
        assert column in time_exit_rows.columns
    components = frames["score_components"].loc[
        frames["score_components"]["hypothesis_id"]
        .astype(str)
        .eq(SECTOR_ROTATION_BUY_ORDINARY_TIME_EXIT_HYPOTHESIS_ID)
    ]
    assert {
        "time_exit_positive_probability_component",
        "expected_time_exit_return_component",
        "time_exit_utility_component",
        "failed_tbs_but_profitable_component",
        "adverse_recovery_penalty",
    }.issubset(set(components["component"].astype(str)))
    selected_feature_payloads = candidate_hypothesis["selected_features"].dropna()
    assert not selected_feature_payloads.empty
    for payload in pd.concat(
        [
            candidate_hypothesis["selected_features"].dropna(),
            time_exit_hypothesis["selected_features"].dropna(),
        ],
        ignore_index=True,
    ):
        selected = json.loads(payload)
        assert not any(str(feature).startswith("label_") for feature in selected)
        assert not any("time_exit" in str(feature) for feature in selected)
        assert not any("utility" in str(feature) for feature in selected)
        assert not any("profitable_despite_failed_tbs" in str(feature) for feature in selected)
        assert not any("early_adverse_recovery" in str(feature) for feature in selected)
    assert model_artifact.read_text(encoding="utf-8") == before_model_artifact
    assert modeling_path.read_bytes() == before_modeling


def test_analog_robustness_classifies_concentrated_top10_decay_as_artifact() -> None:
    frames = historical_analog_robustness_frames_from_analogs(_concentrated_decay_analogs())
    robustness = frames["analog_robustness"].iloc[0]
    flags = set(frames["analog_caution_flags"]["caution_flag"].astype(str))

    assert robustness["original_analog_support_label"] == "SUPPORTIVE"
    assert robustness["robust_analog_support_label"] == "CONCENTRATION_ARTIFACT"
    assert bool(robustness["analog_evidence_too_concentrated"]) is True
    assert "same_symbol_concentration" in flags
    assert "same_year_concentration" in flags
    assert "analogs_mostly_same_event_cluster" in flags
    assert "support_decays_top25" in flags
    assert "support_decays_top50" in flags
    assert "tbs_support_decay" in flags
    assert "high_mae_tail_risk" in flags


def test_analog_robustness_classifies_diversified_depth_support_as_robust() -> None:
    frames = historical_analog_robustness_frames_from_analogs(_diversified_supportive_analogs())
    robustness = frames["analog_robustness"].iloc[0]
    flags = set(frames["analog_caution_flags"]["caution_flag"].astype(str))

    assert robustness["original_analog_support_label"] == "SUPPORTIVE"
    assert robustness["robust_analog_support_label"] == "ROBUST_SUPPORT"
    assert bool(robustness["analog_evidence_usable_for_research"]) is True
    assert "same_symbol_concentration" not in flags
    assert "same_year_concentration" not in flags


def test_analog_robustness_classifies_diversified_depth_decay() -> None:
    frames = historical_analog_robustness_frames_from_analogs(_diversified_decay_analogs())
    robustness = frames["analog_robustness"].iloc[0]
    flags = set(frames["analog_caution_flags"]["caution_flag"].astype(str))

    assert robustness["original_analog_support_label"] == "SUPPORTIVE"
    assert robustness["robust_analog_support_label"] == "DECAYS_WITH_DEPTH"
    assert "support_decays_top25" in flags
    assert "support_decays_top50" in flags
    assert "same_symbol_concentration" not in flags


def test_analog_robustness_low_count_is_insufficient() -> None:
    frame = _diversified_supportive_analogs().head(5)
    frames = historical_analog_robustness_frames_from_analogs(frame)
    robustness = frames["analog_robustness"].iloc[0]
    flags = set(frames["analog_caution_flags"]["caution_flag"].astype(str))

    assert robustness["robust_analog_support_label"] == "INSUFFICIENT_ANALOGS"
    assert "low_analog_count" in flags


def test_signal_discovery_persists_candidates_rejections_scores_and_analogs(
    signal_discovery_root: tuple[Path, Path],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root, config = signal_discovery_root
    before_db = (root / "state" / "engine.sqlite3").read_bytes()
    model_artifact = root / "artifacts" / "models" / "model.joblib"
    model_artifact.parent.mkdir(parents=True, exist_ok=True)
    model_artifact.write_text("model-artifact-do-not-touch", encoding="utf-8")
    before_model_artifact = model_artifact.read_text(encoding="utf-8")

    def fail_download(*_: object, **__: object) -> None:
        raise AssertionError("Signal discovery attempted an FMP request")

    monkeypatch.setattr("swing_rsi.data.loader.download_daily", fail_download)
    run = run_signal_discovery(root, config_path=config)
    frames = load_signal_discovery_frames(root)

    assert run.metadata["schema_version"] == SIGNAL_DISCOVERY_SCHEMA_VERSION
    assert run.metadata["generation_type"] == SIGNAL_DISCOVERY_GENERATION_TYPE
    assert not frames["candidates"].empty
    assert not frames["selected_candidates"].empty
    assert not frames["rejected"].empty
    assert frames["rejected"]["rejection_reason"].astype(str).str.len().gt(0).all()
    assert not frames["score_components"].empty
    assert set(frames["candidates"]["candidate_status"]) <= {
        "SHADOW_ONLY",
        "RESEARCH_ONLY",
        "REJECTED_BY_POLICY",
        "REJECTED_BY_OOD",
    }
    assert "LIVE_ACTIONABLE" not in set(frames["candidates"]["candidate_status"])
    assert frames["selected_candidates"]["footprint_summary"].astype(str).str.len().gt(0).all()
    assert not frames["historical_analogs"].empty
    assert frames["historical_analogs"]["outcome_labels_used_for_explanation_only"].eq(True).all()
    candidates_by_signal = frames["candidates"].set_index("signal_id")
    for _, analog in frames["historical_analogs"].iterrows():
        candidate_date = pd.Timestamp(candidates_by_signal.loc[analog["signal_id"], "as_of_date"])
        assert pd.Timestamp(analog["analog_date"]) < candidate_date
    assert not frames["footprint_evidence"].empty
    assert "Top conflict" in set(frames["footprint_evidence"]["Category"])
    assert "Residual / unexplained" in set(frames["footprint_evidence"]["Category"])
    calibration_summary = frames["calibration_summary"]
    distributions = frames["probability_distributions"]
    thresholds = frames["diagnostic_thresholds"]
    buckets = frames["probability_buckets"]
    row_level = frames["row_level_calibration_audit"]
    manifest = frames["calibration_artifact_manifest"]
    assert not calibration_summary.empty
    assert not distributions.empty
    assert not thresholds.empty
    assert not buckets.empty
    assert not row_level.empty
    assert not manifest.empty
    assert (
        calibration_summary["schema_version"].eq(MULTI_ANGLE_CALIBRATION_AUDIT_SCHEMA_VERSION).all()
    )
    evaluated = set(
        frames["hypotheses"]
        .loc[
            frames["hypotheses"]["family"].astype(str).str.len().gt(0), ["hypothesis_id", "family"]
        ]
        .itertuples(index=False, name=None)
    )
    audited = set(
        calibration_summary.loc[
            calibration_summary["model_family"].astype(str).str.len().gt(0),
            ["hypothesis_id", "model_family"],
        ].itertuples(index=False, name=None)
    )
    assert evaluated.issubset(audited)
    assert {"raw_tbs", "calibrated_tbs"}.issubset(set(distributions["probability_type"]))
    assert {
        "count_ge_030",
        "count_ge_035",
        "count_ge_040",
        "count_ge_045",
        "count_ge_050",
        "count_ge_055",
        "count_ge_060",
    }.issubset(distributions.columns)
    assert set(CALIBRATION_DIAGNOSTIC_THRESHOLDS).issubset(set(thresholds["threshold"].round(2)))
    assert thresholds["diagnostic_only"].eq(True).all()
    assert thresholds["production_target_before_stop_threshold"].eq(0.10).all()
    assert set(row_level["schema_version"]) == {MULTI_ANGLE_CALIBRATION_AUDIT_SCHEMA_VERSION}
    assert "TBS_label" in row_level.columns
    assert row_level["calibration_artifact_hash"].astype(str).str.len().gt(0).all()
    assert (
        pd.to_datetime(row_level["Date"])
        .le(pd.to_datetime(calibration_summary["calibration_end_date"]).max())
        .all()
    )
    assert (
        pd.to_datetime(row_level["Date"])
        .lt(pd.to_datetime(frames["candidates"]["as_of_date"]).max())
        .all()
    )
    selected_feature_payloads = frames["hypotheses"]["selected_features"].dropna()
    assert not selected_feature_payloads.empty
    for payload in selected_feature_payloads:
        selected = json.loads(payload)
        assert not any(str(feature).startswith("label_") for feature in selected)
    assert (root / "state" / "engine.sqlite3").read_bytes() == before_db
    assert model_artifact.read_text(encoding="utf-8") == before_model_artifact


def test_signal_discovery_persists_insufficient_calibration_evidence(tmp_path: Path) -> None:
    root = tmp_path / "insufficient-calibration"
    _write_universe(root)
    _write_feature_data(root)
    config = _write_config(root)
    text = config.read_text(encoding="utf-8")
    text = text.replace("minimum_training_samples: 120", "minimum_training_samples: 20")
    text = text.replace("minimum_calibration_samples: 24", "minimum_calibration_samples: 200")
    text = text.replace("minimum_holdout_samples: 24", "minimum_holdout_samples: 20")
    config.write_text(text, encoding="utf-8")

    run_signal_discovery(root, config_path=config)
    frames = load_signal_discovery_frames(root)
    summary = frames["calibration_summary"]

    assert not summary.empty
    assert set(summary["status"]) == {"INSUFFICIENT_CALIBRATION_EVIDENCE"}
    assert set(summary["reason"]) == {"minimum_split_samples_failed"}
    assert summary["calibration_row_count"].lt(200).all()
    assert frames["probability_distributions"].empty
    assert frames["row_level_calibration_audit"].empty


def test_signal_discovery_persists_no_signal_rows(tmp_path: Path) -> None:
    root = tmp_path / "strict"
    _write_universe(root)
    _write_feature_data(root)
    config = _write_config(
        root,
        signal_score_threshold=0.99,
        probability_threshold=0.99,
        candidate_cap=10,
    )

    run_signal_discovery(root, config_path=config)
    frames = load_signal_discovery_frames(root)

    assert not frames["no_signal"].empty
    assert frames["no_signal"]["decision"].eq("NO_SIGNAL").all()
    assert frames["no_signal"]["no_signal_reason"].astype(str).str.len().gt(0).all()


def test_signal_discovery_export_and_dashboard_workbook_sheets(
    signal_discovery_root: tuple[Path, Path],
    tmp_path: Path,
) -> None:
    root, config = signal_discovery_root
    run_signal_discovery(root, config_path=config)

    output_dir = tmp_path / "exports"
    written = export_signal_discovery_generation(root, generation="latest", output=output_dir)
    written_names = {path.name for path in written}
    frames = load_signal_discovery_frames(root)
    blocked_frames = signal_discovery_blocked_analog_frames(root)
    robustness_frames = signal_discovery_analog_robustness_frames(root)
    workbook = openpyxl.load_workbook(
        BytesIO(
            to_xlsx_bytes(
                {
                    "signal_discovery_summary": frames["summary"],
                    "hypothesis_summary": frames["hypotheses"],
                    "candidate_rows": frames["candidates"],
                    "no_signal_rows": frames["no_signal"],
                    "rejected_rows": frames["rejected"],
                    "footprint_evidence": frames["footprint_evidence"],
                    "historical_analogs": frames["historical_analogs"],
                    "blocked_row_analogs": blocked_frames["blocked_row_analogs"],
                    "blocked_row_analog_summary": blocked_frames["blocked_row_analog_summary"],
                    "analog_robustness": robustness_frames["analog_robustness"],
                    "analog_robustness_summary": robustness_frames["analog_robustness_summary"],
                    "analog_depth_comparison": robustness_frames["analog_depth_comparison"],
                    "analog_caution_flags": robustness_frames["analog_caution_flags"],
                    "calibration_summary": frames["calibration_summary"],
                    "probability_distributions": frames["probability_distributions"],
                    "probability_buckets": frames["probability_buckets"],
                    "diagnostic_thresholds": frames["diagnostic_thresholds"],
                    "row_level_calibration_audit": frames["row_level_calibration_audit"],
                    "calibration_artifact_manifest": frames["calibration_artifact_manifest"],
                    "target_stop_policy_registry": frames["target_stop_policy_registry"],
                    "calibration_selection": frames["calibration_selection"],
                    "baseline_vs_candidate": frames[
                        "sector_rotation_buy_ordinary_policy_comparison"
                    ],
                    "derived_outcomes": frames["derived_policy_outcomes"],
                    "signal_policy_comparison": frames["signal_discovery_policy_comparison"],
                    "time_exit_labels": frames["time_exit_utility_labels"],
                    "time_exit_calibration": frames["time_exit_utility_calibration_summary"],
                    "time_exit_signal_rows": frames["time_exit_utility_signal_rows"],
                    "time_exit_policy_comparison": frames["time_exit_utility_policy_comparison"],
                }
            )
        )
    )

    assert "candidates.csv" in written_names
    assert "blocked_row_analogs.csv" in written_names
    assert "blocked_row_analog_summary.csv" in written_names
    assert "analog_robustness.csv" in written_names
    assert "analog_robustness_summary.csv" in written_names
    assert "analog_depth_comparison.csv" in written_names
    assert "analog_caution_flags.csv" in written_names
    assert "calibration_summary.csv" in written_names
    assert "calibration_summary.json" in written_names
    assert "probability_distributions.csv" in written_names
    assert "probability_buckets.csv" in written_names
    assert "diagnostic_thresholds.csv" in written_names
    assert "row_level_calibration_audit.csv" in written_names
    assert "row_level_calibration_audit.parquet" in written_names
    assert "calibration_artifact_manifest.json" in written_names
    assert "target_stop_policy_registry.csv" in written_names
    assert "target_stop_policy_registry.json" in written_names
    assert "calibration_selection.csv" in written_names
    assert "sector_rotation_buy_ordinary_policy_comparison.csv" in written_names
    assert "derived_policy_outcomes.csv" in written_names
    assert "signal_discovery_policy_comparison.csv" in written_names
    assert "time_exit_utility_labels.csv" in written_names
    assert "time_exit_utility_calibration_summary.csv" in written_names
    assert "time_exit_utility_signal_rows.csv" in written_names
    assert "time_exit_utility_policy_comparison.csv" in written_names
    assert "metadata.json" in written_names
    assert {
        "signal_discovery_summary",
        "hypothesis_summary",
        "candidate_rows",
        "no_signal_rows",
        "rejected_rows",
        "footprint_evidence",
        "historical_analogs",
        "blocked_row_analogs",
        "blocked_row_analog_summary",
        "analog_robustness",
        "analog_robustness_summary",
        "analog_depth_comparison",
        "analog_caution_flags",
        "calibration_summary",
        "probability_distributions",
        "probability_buckets",
        "diagnostic_thresholds",
        "row_level_calibration_audit",
        "calibration_artifact_manifest",
        "target_stop_policy_registry",
        "calibration_selection",
        "baseline_vs_candidate",
        "derived_outcomes",
        "signal_policy_comparison",
        "time_exit_labels",
        "time_exit_calibration",
        "time_exit_signal_rows",
        "time_exit_policy_comparison",
    }.issubset(set(workbook.sheetnames))


def test_signal_discovery_blocker_report_summarizes_reasons_and_groups(
    signal_discovery_root: tuple[Path, Path],
    tmp_path: Path,
) -> None:
    root, config = signal_discovery_root
    run_signal_discovery(root, config_path=config)
    report = signal_discovery_blocker_report_frames(root)

    assert not report["summary"].empty
    assert not report["blocker_rows"].empty
    assert not report["by_reason"].empty
    assert not report["by_hypothesis"].empty
    assert not report["by_archetype"].empty
    assert not report["by_ticker"].empty
    assert not report["by_scope"].empty
    assert {"NO_SIGNAL", "REJECTED_BY_POLICY"}.intersection(
        set(report["blocker_rows"]["decision"].astype(str))
    )
    assert "blocker_reason" in report["by_reason"].columns
    assert "hypothesis_id" in report["by_hypothesis"].columns
    assert "archetype" in report["by_archetype"].columns
    assert "ticker" in report["by_ticker"].columns
    assert "product_class_scope" in report["by_scope"].columns
    assert "signal_id" in report["blocker_rows"].columns
    assert "model_id" in report["blocker_rows"].columns
    assert report["blocker_rows"]["model_id"].astype(str).str.len().gt(0).all()
    assert report["summary"]["blocker_rows"].iloc[0] == len(report["blocker_rows"])

    output_dir = tmp_path / "blocker_export"
    written = export_signal_discovery_blocker_report(
        root,
        generation="latest",
        output=output_dir,
    )
    names = {path.name for path in written}
    workbook = openpyxl.load_workbook(BytesIO(to_xlsx_bytes(report)))

    assert "by_reason.csv" in names
    assert "blocker_rows.csv" in names
    assert "metadata.json" in names
    assert {
        "summary",
        "metadata",
        "blocker_rows",
        "by_reason",
        "by_hypothesis",
        "by_archetype",
        "by_ticker",
        "by_scope",
    }.issubset(set(workbook.sheetnames))


def test_blocked_row_analogs_are_chronological_scope_labeled_and_explanatory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "blocked"
    _write_blocked_analog_fixture(root)
    artifact = root / "artifacts" / "models" / "model.joblib"
    artifact.parent.mkdir(parents=True, exist_ok=True)
    artifact.write_text("do-not-touch", encoding="utf-8")
    operational = tmp_path / "operational"
    operational.mkdir()
    marker = operational / "state-marker.txt"
    marker.write_text("operational-untouched", encoding="utf-8")

    def fail_download(*_: object, **__: object) -> None:
        raise AssertionError("Blocked-row analog diagnostics attempted an FMP request")

    monkeypatch.setattr("swing_rsi.data.loader.download_daily", fail_download)
    before_artifact = artifact.read_text(encoding="utf-8")
    before_marker = marker.read_text(encoding="utf-8")

    frames = signal_discovery_blocked_analog_frames(root, analog_count=4)
    robustness_frames = signal_discovery_analog_robustness_frames(root)
    analogs = frames["blocked_row_analogs"]
    summary = frames["blocked_row_analog_summary"]
    robustness = robustness_frames["analog_robustness"]
    depth = robustness_frames["analog_depth_comparison"]
    tza_analogs = analogs.loc[analogs["target_ticker"].eq("TZA")]
    tza_summary = summary.loc[summary["target_ticker"].eq("TZA")].iloc[0]

    assert len(tza_analogs) == 4
    assert pd.to_datetime(tza_analogs["analog_date"]).lt(pd.Timestamp("2026-06-26")).all()
    assert "2026-06-26" not in set(tza_analogs["analog_date"].astype(str))
    assert "2026-06-29" not in set(analogs["analog_date"].astype(str))
    assert tza_analogs.iloc[0]["analog_ticker"] == "TZA"
    assert tza_analogs.head(2)["same_product_scope"].eq(True).all()
    assert "cross_scope_fallback" in set(tza_analogs["analog_pool"])
    assert tza_analogs["outcome_labels_used_for_explanation_only"].eq(True).all()
    assert tza_summary["same_scope_analog_count"] == 2
    assert math.isclose(float(tza_summary["average_forward_return"]), 0.01)
    assert math.isclose(float(tza_summary["average_MFE"]), 0.05)
    assert math.isclose(float(tza_summary["average_MAE"]), -0.03)
    assert math.isclose(float(tza_summary["target_before_stop_hit_rate"]), 0.5)
    assert tza_summary["analog_support_label"] == "MIXED"
    assert "explanatory only" in str(tza_summary["analog_footprint_summary"])
    assert set(summary.loc[summary["target_ticker"].eq("SOXS"), "key_caution"]) == {
        "OOD target row"
    }
    assert set(summary["target_status"]) <= {"RESEARCH_ONLY", "REJECTED_BY_OOD"}
    assert set(robustness["target_status"]) <= {"RESEARCH_ONLY", "REJECTED_BY_OOD"}
    assert set(robustness["robust_analog_support_label"]).issubset(
        {
            "INSUFFICIENT_ANALOGS",
            "MIXED_SUPPORT",
            "CONCENTRATION_ARTIFACT",
            "DECAYS_WITH_DEPTH",
            "ROBUST_SUPPORT",
            "SUPPORTIVE_BUT_CONCENTRATED",
            "WEAK_SUPPORT",
        }
    )
    assert pd.to_datetime(analogs["analog_date"]).lt(pd.Timestamp("2026-06-26")).all()
    assert depth["target_signal_id"].isin(set(robustness["target_signal_id"])).all()
    assert artifact.read_text(encoding="utf-8") == before_artifact
    assert marker.read_text(encoding="utf-8") == before_marker


def test_blocked_row_analog_exports_include_csv_and_workbook_sheets(tmp_path: Path) -> None:
    root = tmp_path / "blocked_export"
    _write_blocked_analog_fixture(root)

    output = tmp_path / "export"
    written = export_signal_discovery_generation(root, generation="latest", output=output)
    names = {path.name for path in written}
    frames = signal_discovery_blocked_analog_frames(root)
    robustness_frames = signal_discovery_analog_robustness_frames(root)
    workbook = openpyxl.load_workbook(
        BytesIO(
            to_xlsx_bytes(
                {
                    "blocked_row_analogs": frames["blocked_row_analogs"],
                    "blocked_row_analog_summary": frames["blocked_row_analog_summary"],
                    "analog_robustness": robustness_frames["analog_robustness"],
                    "analog_robustness_summary": robustness_frames["analog_robustness_summary"],
                    "analog_depth_comparison": robustness_frames["analog_depth_comparison"],
                    "analog_caution_flags": robustness_frames["analog_caution_flags"],
                }
            )
        )
    )

    assert "blocked_row_analogs.csv" in names
    assert "blocked_row_analog_summary.csv" in names
    assert "analog_robustness.csv" in names
    assert "analog_robustness_summary.csv" in names
    assert "analog_depth_comparison.csv" in names
    assert "analog_caution_flags.csv" in names
    assert {
        "blocked_row_analogs",
        "blocked_row_analog_summary",
        "analog_robustness",
        "analog_robustness_summary",
        "analog_depth_comparison",
        "analog_caution_flags",
    }.issubset(set(workbook.sheetnames))


def test_signal_discovery_leaves_operational_like_path_untouched(tmp_path: Path) -> None:
    operational = tmp_path / "operational"
    operational.mkdir()
    marker = operational / "state-marker.txt"
    marker.write_text("do-not-touch", encoding="utf-8")
    root = tmp_path / "dev"
    _write_universe(root)
    _write_feature_data(root)
    config = _write_config(root)

    before = marker.read_text(encoding="utf-8")
    run_signal_discovery(root, config_path=config)

    assert marker.read_text(encoding="utf-8") == before
